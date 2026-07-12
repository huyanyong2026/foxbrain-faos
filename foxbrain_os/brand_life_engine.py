"""Brand Life Engine: evidence-backed brand knowledge understanding layer."""

from __future__ import annotations

import hashlib
import json
import re
import time
from pathlib import Path


LIFE_DIMENSIONS = (
    "brand_identity",
    "brand_story",
    "brand_philosophy",
    "product_system",
    "store_system",
    "people_system",
    "training_system",
    "customer_system",
    "operating_rules",
    "future_plan",
)

KNOWLEDGE_CLASSES = {
    "brand_life": "品牌生命",
    "people_life": "人才生命",
    "store_life": "门店生命",
    "training_knowledge": "培训知识",
}

PUBLIC_ROLES = ("employee", "store_manager", "purchasing", "finance", "boss", "admin")
SENSITIVE_ROLES = ("boss", "admin")


def extract_brand_document(source_path):
    path = Path(source_path)
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(str(path))
            text = "\n".join(f"[第{index}页]\n{page.extract_text() or ''}" for index, page in enumerate(reader.pages, 1))
            return {"text": text[:400000], "status": "parsed" if text.strip() else "needs_ocr", "error": ""}
        if ext == ".docx":
            import docx
            document = docx.Document(str(path))
            text = "\n".join(f"[段落{index}] {paragraph.text}" for index, paragraph in enumerate(document.paragraphs, 1) if paragraph.text.strip())
            return {"text": text[:400000], "status": "parsed" if text else "needs_review", "error": ""}
        if ext in (".xlsx", ".xls"):
            lines = []
            try:
                import openpyxl
                workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
                for sheet in workbook.worksheets:
                    lines.append(f"[工作表:{sheet.title}]")
                    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        values = "\t".join("" if value is None else str(value) for value in row)
                        if values.strip():
                            lines.append(f"[第{row_number}行] {values}")
            except Exception:
                raw = path.read_bytes()
                decoded = ""
                for encoding in ("utf-8-sig", "gb18030", "gbk"):
                    try:
                        decoded = raw.decode(encoding)
                        break
                    except UnicodeDecodeError:
                        continue
                if decoded:
                    lines.append("[工作表:SAP文本导出]")
                    lines.extend(f"[第{index}行] {line}" for index, line in enumerate(decoded.splitlines(), 1))
            text = "\n".join(lines)
            return {"text": text[:400000], "status": "parsed" if text else "needs_review", "error": ""}
        if ext in (".jpg", ".jpeg", ".png", ".webp"):
            from PIL import Image
            import pytesseract
            text = pytesseract.image_to_string(Image.open(path), lang="chi_sim+eng")
            return {"text": text[:400000], "status": "parsed" if text.strip() else "needs_ocr", "error": ""}
        if ext in (".txt", ".md", ".csv", ".tsv"):
            text = path.read_text(encoding="utf-8", errors="ignore")
            return {"text": text[:400000], "status": "parsed", "error": ""}
        return {"text": "", "status": "unsupported", "error": "暂不支持该格式"}
    except Exception as exc:
        status = "needs_ocr" if ext in (".pdf", ".jpg", ".jpeg", ".png", ".webp") else "needs_review"
        return {"text": "", "status": status, "error": str(exc)[:500]}


def ensure_brand_life_schema(conn):
    conn.execute(
        """create table if not exists brand_life_profiles(
        id integer primary key autoincrement, brand_code text not null unique,
        brand_name text not null, brand_identity text, brand_story text,
        brand_philosophy text, product_system text, store_system text,
        people_system text, training_system text, customer_system text,
        operating_rules text, future_plan text, status text not null default 'active',
        created_at integer not null, updated_at integer not null)"""
    )
    conn.execute(
        """create table if not exists brand_knowledge_vault_items(
        id integer primary key autoincrement, brand_id integer not null,
        document_id integer, title text not null, source_filename text,
        source_path text, content_hash text not null, file_type text,
        confidentiality text not null default 'sensitive',
        processing_status text not null default 'pending', ai_summary text,
        tags text, extracted_text text, processing_error text,
        created_by integer, created_at integer not null, updated_at integer not null,
        unique(brand_id,content_hash))"""
    )
    conn.execute(
        """create table if not exists brand_vault_item_categories(
        id integer primary key autoincrement, vault_item_id integer not null,
        category text not null, confidence real not null default 1,
        evidence text, confirmed integer not null default 0,
        confirmed_by integer, created_at integer not null,
        unique(vault_item_id,category))"""
    )
    conn.execute(
        """create table if not exists brand_knowledge_cards(
        id integer primary key autoincrement, brand_id integer not null,
        vault_item_id integer not null, category text not null, title text not null,
        content text not null, evidence_location text not null,
        confidentiality text not null default 'sensitive',
        review_status text not null default 'draft', created_by integer,
        created_at integer not null, updated_at integer not null)"""
    )
    conn.execute(
        """create table if not exists brand_life_permissions(
        id integer primary key autoincrement, brand_id integer not null,
        role text not null, knowledge_scope text not null,
        permission text not null default 'read', created_at integer not null,
        unique(brand_id,role,knowledge_scope))"""
    )
    for statement in (
        "create index if not exists idx_brand_vault_brand on brand_knowledge_vault_items(brand_id,confidentiality,processing_status)",
        "create index if not exists idx_brand_cards_brand on brand_knowledge_cards(brand_id,category,review_status)",
        "create index if not exists idx_brand_categories_item on brand_vault_item_categories(vault_item_id,category)",
    ):
        conn.execute(statement)


def seed_kailas(conn):
    ensure_brand_life_schema(conn)
    now = int(time.time())
    row = conn.execute("select id from brand_life_profiles where brand_code='KAILAS'").fetchone()
    if row:
        brand_id = row[0]
    else:
        brand_id = conn.execute(
            """insert into brand_life_profiles(
            brand_code,brand_name,brand_identity,brand_philosophy,status,created_at,updated_at)
            values(?,?,?,?,?,?,?)""",
            (
                "KAILAS",
                "KAILAS 凯乐石",
                "中国专业户外品牌；详细身份资料等待品牌文件持续补充。",
                "专业、安全与山地运动体验；正式表述以品牌授权资料为准。",
                "active",
                now,
                now,
            ),
        ).lastrowid
    for role in PUBLIC_ROLES:
        conn.execute(
            "insert or ignore into brand_life_permissions(brand_id,role,knowledge_scope,permission,created_at) values(?,?,?,?,?)",
            (brand_id, role, "public", "read", now),
        )
    for role in SENSITIVE_ROLES:
        conn.execute(
            "insert or ignore into brand_life_permissions(brand_id,role,knowledge_scope,permission,created_at) values(?,?,?,?,?)",
            (brand_id, role, "sensitive", "read", now),
        )
    return brand_id


def classify_brand_document(filename, text):
    haystack = ((filename or "") + "\n" + (text or "")).lower()
    rules = {
        "brand_life": ("kailas", "凯乐石", "品牌", "理念"),
        "people_life": ("薪酬", "福利", "员工", "人才", "岗位", "奖金", "提成"),
        "store_life": ("终端", "门店", "店长", "加盟", "零售"),
        "training_knowledge": ("培训", "考核", "学习", "课程", "标准", "晋升"),
    }
    result = []
    for category, words in rules.items():
        hits = [word for word in words if word in haystack]
        if hits:
            result.append(
                {
                    "category": category,
                    "confidence": min(0.98, 0.65 + 0.08 * len(hits)),
                    "evidence": "、".join(hits[:5]),
                }
            )
    return result or [{"category": "brand_life", "confidence": 0.5, "evidence": "品牌资料待人工复核"}]


def split_evidence(text, size=900):
    clean = re.sub(r"\s+", " ", text or "").strip()
    return [clean[index:index + size] for index in range(0, len(clean), size) if clean[index:index + size]]


def build_knowledge_cards(filename, text, categories, confidentiality="sensitive"):
    chunks = split_evidence(text)
    if not chunks:
        return []
    cards = []
    for index, chunk in enumerate(chunks[:12], 1):
        category = categories[(index - 1) % len(categories)]["category"]
        cards.append(
            {
                "category": category,
                "title": f"{KNOWLEDGE_CLASSES[category]}知识卡 {index}",
                "content": chunk[:900],
                "evidence_location": f"{filename} / 文本段落 {index}",
                "confidentiality": confidentiality,
                "review_status": "draft",
            }
        )
    return cards


def register_brand_document(
    conn,
    brand_id,
    filename,
    source_path,
    extracted_text,
    document_id=None,
    created_by=None,
    confidentiality="sensitive",
    processing_error="",
):
    ensure_brand_life_schema(conn)
    path = Path(source_path) if source_path else None
    content_hash = hashlib.sha256(path.read_bytes()).hexdigest() if path and path.exists() else hashlib.sha256((filename + extracted_text).encode("utf-8")).hexdigest()
    categories = classify_brand_document(filename, extracted_text)
    cards = build_knowledge_cards(filename, extracted_text, categories, confidentiality)
    now = int(time.time())
    existing = conn.execute(
        "select id from brand_knowledge_vault_items where brand_id=? and content_hash=?",
        (brand_id, content_hash),
    ).fetchone()
    if existing:
        return {"ok": True, "duplicate": True, "vault_item_id": existing[0], "categories": categories, "cards": 0}
    summary = re.sub(r"\s+", " ", extracted_text or "").strip()[:300]
    status = "processed" if extracted_text else ("needs_ocr" if Path(filename).suffix.lower() in (".jpg", ".jpeg", ".png", ".webp") else "needs_review")
    cur = conn.execute(
        """insert into brand_knowledge_vault_items(
        brand_id,document_id,title,source_filename,source_path,content_hash,file_type,
        confidentiality,processing_status,ai_summary,tags,extracted_text,processing_error,
        created_by,created_at,updated_at) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (brand_id, document_id, filename, filename, str(source_path or ""), content_hash,
         Path(filename).suffix.lower().lstrip("."), confidentiality, status, summary,
         ",".join(KNOWLEDGE_CLASSES[item["category"]] for item in categories),
         extracted_text, processing_error, created_by, now, now),
    )
    vault_item_id = cur.lastrowid
    for item in categories:
        conn.execute(
            "insert into brand_vault_item_categories(vault_item_id,category,confidence,evidence,confirmed,created_at) values(?,?,?,?,?,?)",
            (vault_item_id, item["category"], item["confidence"], item["evidence"], 0, now),
        )
    for card in cards:
        conn.execute(
            """insert into brand_knowledge_cards(
            brand_id,vault_item_id,category,title,content,evidence_location,
            confidentiality,review_status,created_by,created_at,updated_at)
            values(?,?,?,?,?,?,?,?,?,?,?)""",
            (brand_id, vault_item_id, card["category"], card["title"], card["content"],
             card["evidence_location"], card["confidentiality"], card["review_status"],
             created_by, now, now),
        )
    return {"ok": True, "duplicate": False, "vault_item_id": vault_item_id, "categories": categories, "cards": len(cards), "status": status}


def can_read_brand_knowledge(role, confidentiality):
    if confidentiality == "public":
        return role in PUBLIC_ROLES
    return role in SENSITIVE_ROLES


def brand_life_payload(conn, brand_code="KAILAS", role="employee"):
    brand = conn.execute("select * from brand_life_profiles where brand_code=?", (brand_code,)).fetchone()
    if not brand:
        return {"ok": False, "message": "品牌尚未建立"}
    brand_id = brand["id"] if hasattr(brand, "keys") else brand[0]
    allowed = ["public"] + (["sensitive"] if role in SENSITIVE_ROLES else [])
    placeholders = ",".join("?" for _ in allowed)
    vault = conn.execute(
        f"select * from brand_knowledge_vault_items where brand_id=? and confidentiality in ({placeholders}) order by created_at desc",
        [brand_id] + allowed,
    ).fetchall()
    cards = conn.execute(
        f"select * from brand_knowledge_cards where brand_id=? and confidentiality in ({placeholders}) order by created_at desc limit 50",
        [brand_id] + allowed,
    ).fetchall()
    return {
        "ok": True,
        "brand": dict(brand),
        "vault_items": [dict(row) for row in vault],
        "knowledge_cards": [dict(row) for row in cards],
        "visible_scopes": allowed,
        "dimensions": list(LIFE_DIMENSIONS),
    }
