#!/usr/bin/env python3
import base64
import cgi
import csv
import hashlib
import hmac
import html
import io
import json
import mimetypes
import os
import re
import secrets
import sqlite3
import time
import urllib.request
import uuid
from http import cookies
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse


APP_DIR = "/opt/firefox-portal"
DB = APP_DIR + "/portal.db"
SECRET_FILE = APP_DIR + "/secret.key"
ENV_FILE = APP_DIR + "/portal.env"
SAP_SUMMARY_FILE = "/opt/firefox-sap-sync/latest_summary.json"
UPLOAD_DIR = APP_DIR + "/uploads"
PORT = 8088
LOCK_LIMIT = 5
LOCK_SECONDS = 15 * 60


def U(s):
    return s.encode("ascii").decode("unicode_escape")


def load_env_file():
    if not os.path.exists(ENV_FILE):
        return
    for line in Path(ENV_FILE).read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, val = line.split("=", 1)
        os.environ.setdefault(key.strip(), val.strip())


load_env_file()

T = {
    "brand": U(r"\u706b\u72d0\u72f8 AI \u4f01\u4e1a\u7ecf\u8425\u7cfb\u7edf"),
    "subtitle": U(r"FireFox AI Operating System\uff1aAI + ERP + CRM + OA + \u77e5\u8bc6\u5e93 + BI + \u667a\u80fd\u4f53\u5e73\u53f0\u7684\u7edf\u4e00\u5165\u53e3\u3002"),
    "login": U(r"\u767b\u5f55"),
    "register": U(r"\u65b0\u5458\u5de5\u6ce8\u518c"),
    "logout": U(r"\u9000\u51fa"),
    "email": U(r"\u90ae\u7bb1"),
    "password": U(r"\u5bc6\u7801"),
    "new_password": U(r"\u65b0\u5bc6\u7801"),
    "name": U(r"\u59d3\u540d"),
    "phone": U(r"\u624b\u673a"),
    "store": U(r"\u95e8\u5e97/\u90e8\u95e8"),
    "role": U(r"\u89d2\u8272"),
    "status": U(r"\u72b6\u6001"),
    "action": U(r"\u64cd\u4f5c"),
    "pending": U(r"\u4f60\u7684\u8d26\u53f7\u5df2\u63d0\u4ea4\uff0c\u9700\u7ba1\u7406\u5458\u5ba1\u6838\u901a\u8fc7\u540e\u624d\u80fd\u767b\u5f55\u3002"),
    "bad": U(r"\u8d26\u53f7\u6216\u5bc6\u7801\u4e0d\u6b63\u786e\uff0c\u6216\u8d26\u53f7\u8fd8\u6ca1\u6709\u5ba1\u6838\u901a\u8fc7\u3002"),
    "locked": U(r"\u767b\u5f55\u5931\u8d25\u6b21\u6570\u8fc7\u591a\uff0c\u8d26\u53f7\u5df2\u9501\u5b9a 15 \u5206\u949f\u3002"),
    "duplicate": U(r"\u8fd9\u4e2a\u90ae\u7bb1\u5df2\u7ecf\u6ce8\u518c\u8fc7\u3002"),
    "admin": U(r"\u7cfb\u7edf\u7ba1\u7406"),
    "users": U(r"\u7528\u6237\u7ba1\u7406"),
    "change_password": U(r"\u4fee\u6539\u5bc6\u7801"),
    "save": U(r"\u4fdd\u5b58"),
    "approve": U(r"\u5ba1\u6838\u901a\u8fc7"),
    "disable": U(r"\u7981\u7528"),
    "enable": U(r"\u542f\u7528"),
    "reset": U(r"\u91cd\u7f6e\u5bc6\u7801"),
    "no_permission": U(r"\u6ca1\u6709\u6743\u9650"),
    "password_changed": U(r"\u5bc6\u7801\u5df2\u4fee\u6539\u3002"),
}

ROLES = {
    "boss": U(r"\u8001\u677f"),
    "store_manager": U(r"\u5e97\u957f"),
    "employee": U(r"\u5458\u5de5"),
    "purchasing": U(r"\u91c7\u8d2d"),
    "finance": U(r"\u8d22\u52a1"),
    "admin": U(r"\u7ba1\u7406\u5458"),
}

MODULES = {
    "/overview": (U(r"\u7ecf\u8425\u603b\u89c8"), U(r"\u67e5\u770b\u9500\u552e\u3001\u6bdb\u5229\u3001\u5e93\u5b58\u3001\u98ce\u9669\u548c AI \u7ecf\u8425\u5efa\u8bae\u3002"), ("boss", "admin", "finance", "purchasing")),
    "/stores": (U(r"\u95e8\u5e97\u4e2d\u5fc3"), U(r"\u95e8\u5e97\u6863\u6848\u3001\u7ecf\u8425\u6570\u636e\u3001\u95e8\u5e97\u65f6\u95f4\u8f74\u548c AI \u5206\u6790\u3002"), ("boss", "admin", "store_manager", "finance")),
    "/employees": (U(r"\u5458\u5de5\u4e2d\u5fc3"), U(r"\u5458\u5de5\u6863\u6848\u3001\u9500\u552e\u8868\u73b0\u3001\u57f9\u8bad\u8bb0\u5f55\u3001\u6210\u957f\u65f6\u95f4\u8f74\u3002"), ("boss", "admin", "store_manager", "finance")),
    "/brands": (U(r"\u54c1\u724c\u4e2d\u5fc3"), U(r"\u54c1\u724c\u6863\u6848\u3001\u5408\u4f5c\u6587\u4ef6\u3001\u54c1\u724c\u9500\u552e\u548c AI \u7ecf\u8425\u5efa\u8bae\u3002"), ("boss", "admin", "purchasing", "store_manager")),
    "/products": (U(r"\u4ea7\u54c1\u4e2d\u5fc3"), U(r"\u4ea7\u54c1\u6863\u6848\u3001\u56fe\u7247\u3001\u5e93\u5b58\u3001\u9500\u552e\u6570\u636e\u3001AI \u9500\u552e\u8bdd\u672f\u3002"), ("boss", "admin", "purchasing", "store_manager", "employee")),
    "/suppliers": (U(r"\u4f9b\u5e94\u5546\u4e2d\u5fc3"), U(r"\u4f9b\u5e94\u5546\u6863\u6848\u3001\u5408\u540c\u3001\u91c7\u8d2d\u8bb0\u5f55\u3001\u4ed8\u6b3e\u8bb0\u5f55\u548c AI \u8bc4\u4f30\u3002"), ("boss", "admin", "purchasing", "finance")),
    "/members": (U(r"\u987e\u5ba2/\u4f1a\u5458\u4e2d\u5fc3"), U(r"\u4f1a\u5458\u6863\u6848\u3001\u8d2d\u4e70\u5386\u53f2\u3001\u504f\u597d\u6807\u7b7e\u3001AI \u7ef4\u62a4\u5efa\u8bae\u3002"), ("boss", "admin", "store_manager", "employee")),
    "/finance": (U(r"\u8d22\u52a1\u4e2d\u5fc3"), U(r"\u5bf9\u516c\u5408\u89c4\u8d26\u52a1\u3001\u5bf9\u79c1\u73b0\u91d1\u8d26\u3001\u4ed8\u6b3e\u8bc4\u4f30\u548c\u8d44\u91d1\u8ba1\u5212\u3002"), ("boss", "admin", "finance")),
    "/content": (U(r"\u5185\u5bb9\u53d1\u5e03\u4e2d\u5fc3"), U(r"\u65b0\u5a92\u4f53\u63a8\u5e7f\u3001\u95e8\u5e97\u5185\u5bb9\u3001\u4ea7\u54c1\u7d20\u6750\u548c\u53d1\u5e03\u8ba1\u5212\u3002"), ("boss", "admin", "store_manager", "employee")),
    "/tasks": (U(r"\u4efb\u52a1\u4e2d\u5fc3"), U(r"\u4eca\u65e5\u5f85\u529e\u3001\u95e8\u5e97\u4efb\u52a1\u3001\u5458\u5de5\u8ddf\u8fdb\u548c\u81ea\u52a8\u5316\u4efb\u52a1\u3002"), ("boss", "admin", "store_manager", "employee", "purchasing", "finance")),
}

ARCHIVE_FIELDS = {
    "stores": [U(r"\u95e8\u5e97\u540d\u79f0"), U(r"\u5730\u5740"), U(r"\u9762\u79ef"), U(r"\u5f00\u4e1a\u65f6\u95f4"), U(r"\u79df\u8d41\u5408\u540c"), U(r"\u8425\u4e1a\u989d"), U(r"\u79df\u91d1"), U(r"\u5458\u5de5"), U(r"\u54c1\u724c"), U(r"AI \u5efa\u8bae")],
    "employees": [U(r"\u59d3\u540d"), U(r"\u7167\u7247"), U(r"\u5e74\u9f84"), U(r"\u7535\u8bdd"), U(r"\u5c97\u4f4d"), U(r"\u90e8\u95e8"), U(r"\u5165\u804c\u65e5\u671f"), U(r"\u5de5\u8d44"), U(r"\u7ee9\u6548"), U(r"AI \u8bc4\u4ef7")],
    "brands": [U(r"\u54c1\u724c\u540d\u79f0"), "Logo", U(r"\u54c1\u724c\u4ecb\u7ecd"), U(r"\u8d1f\u8d23\u4eba"), U(r"\u5408\u540c"), U(r"\u9500\u552e"), U(r"\u6bdb\u5229"), U(r"\u5e93\u5b58"), U(r"\u672a\u6765\u89c4\u5212"), U(r"AI \u5206\u6790")],
    "products": [U(r"\u4ea7\u54c1\u540d\u79f0"), "SKU", U(r"\u6761\u7801"), U(r"\u54c1\u724c"), U(r"\u5206\u7c7b"), U(r"\u56fe\u7247"), U(r"\u8bf4\u660e\u4e66"), U(r"\u9500\u552e"), U(r"\u5e93\u5b58"), U(r"AI \u63a8\u8350")],
    "suppliers": [U(r"\u4f9b\u5e94\u5546\u540d\u79f0"), U(r"\u8054\u7cfb\u4eba"), U(r"\u7535\u8bdd"), U(r"\u5fae\u4fe1"), U(r"\u5408\u540c"), U(r"\u4ed8\u6b3e\u65b9\u5f0f"), U(r"\u5386\u53f2\u91c7\u8d2d"), U(r"\u9644\u4ef6"), U(r"AI \u8bc4\u4ef7")],
    "members": [U(r"\u59d3\u540d"), U(r"\u7535\u8bdd"), U(r"\u5fae\u4fe1"), U(r"\u751f\u65e5"), U(r"\u6d88\u8d39\u8bb0\u5f55"), U(r"\u79ef\u5206"), U(r"\u4f1a\u5458\u7b49\u7ea7"), U(r"\u5174\u8da3"), U(r"\u5907\u6ce8"), U(r"AI \u63a8\u8350")],
}
OLD_ROLE_MAP = {
    "leader": "boss",
    "manager": "store_manager",
}
STATUS = {
    "pending": U(r"\u5f85\u5ba1\u6838"),
    "approved": U(r"\u6b63\u5e38"),
    "disabled": U(r"\u5df2\u7981\u7528"),
}


def esc(value):
    return html.escape(str(value or ""))


def money(value):
    try:
        return f"{float(value):,.0f}"
    except Exception:
        return "0"


def pct(value):
    try:
        return f"{float(value):.1f}%"
    except Exception:
        return "0.0%"


def ts():
    return int(time.time())


def dt(value):
    try:
        return time.strftime("%Y-%m-%d %H:%M", time.localtime(int(value or 0)))
    except Exception:
        return ""


def module_key(path_or_key):
    return str(path_or_key or "").strip().lstrip("/")


def module_title(key):
    data = MODULES.get("/" + module_key(key))
    return data[0] if data else module_key(key)


def classify_text(text, fallback="01_公司制度"):
    text = (text or "").lower()
    rules = [
        ("02_产品资料", ["产品", "货号", "尺码", "面料", "价格", "gtx", "gore", "kailas"]),
        ("03_品牌资料", ["品牌", "logo", "合同", "mammut", "osprey", "salomon"]),
        ("04_销售话术", ["话术", "顾客", "推荐", "成交", "异议"]),
        ("06_门店经营", ["门店", "店长", "销售", "日报", "任务"]),
        ("07_库存采购", ["库存", "采购", "供应商", "订货", "补货"]),
        ("09_财务经营", ["财务", "付款", "现金", "社保", "工资", "公积金"]),
        ("14_AI经营日报", ["ai", "晨报", "经营分析", "建议"]),
    ]
    for category, words in rules:
        if any(w.lower() in text for w in words):
            return category
    return fallback


def summarize_text(text, limit=220):
    clean = re.sub(r"\s+", " ", text or "").strip()
    if not clean:
        return U(r"\u6682\u65e0\u53ef\u63d0\u53d6\u6587\u672c\uff0c\u5df2\u5148\u4fdd\u5b58\u539f\u6587\u4ef6\u3002")
    return clean[:limit] + ("..." if len(clean) > limit else "")


def extract_tags(text):
    category = classify_text(text)
    tags = [category.split("_", 1)[-1]]
    for word in [U(r"\u9500\u552e"), U(r"\u5e93\u5b58"), U(r"\u91c7\u8d2d"), U(r"\u4f1a\u5458"), U(r"\u95e8\u5e97"), "SAP B1", "AI"]:
        if word.lower() in (text or "").lower():
            tags.append(word)
    return ",".join(dict.fromkeys(tags))


def extract_file_text(path, filename):
    ext = Path(filename).suffix.lower()
    try:
        if ext in (".txt", ".md", ".csv", ".json"):
            return Path(path).read_text(encoding="utf-8", errors="ignore")[:200000]
        if ext == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(path)
                return "\n".join((p.extract_text() or "") for p in reader.pages)[:200000]
            except Exception:
                return ""
        if ext == ".docx":
            try:
                import docx
                doc = docx.Document(path)
                return "\n".join(p.text for p in doc.paragraphs)[:200000]
            except Exception:
                return ""
        if ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                lines = []
                for ws in wb.worksheets[:3]:
                    for row in ws.iter_rows(max_row=200, values_only=True):
                        lines.append("\t".join("" if v is None else str(v) for v in row))
                return "\n".join(lines)[:200000]
            except Exception:
                return ""
    except Exception:
        return ""
    return ""


def fetch_url_text(url):
    if not url.lower().startswith(("http://", "https://")):
        return ""
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 FoxBrain"})
        with urllib.request.urlopen(req, timeout=10) as resp:
            raw = resp.read(500000)
        text = raw.decode("utf-8", errors="ignore")
        text = re.sub(r"(?is)<script.*?</script>|<style.*?</style>", " ", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        return re.sub(r"\s+", " ", html.unescape(text)).strip()[:200000]
    except Exception:
        return ""


def load_summary():
    fallback = {
        "data_date": U(r"\u6d4b\u8bd5\u6570\u636e"),
        "yesterday_sales": 80539.66,
        "yesterday_gross_profit": 23848.03,
        "yesterday_gross_margin": 29.6,
        "month_sales": 118059.57,
        "month_gross_profit": 35385.52,
        "month_target": 900000,
        "completion_rate": 13.1,
        "inventory_amount": 0,
        "risk_count": 0,
        "top_stores": [
            {"store": "8001", "sales": 43866.40, "gross_profit": 12642.40},
            {"store": "8002", "sales": 36844.00, "gross_profit": 9200.43},
            {"store": "8003", "sales": 27678.17, "gross_profit": 9195.06},
        ],
        "ai_suggestions": [
            U(r"\u4eca\u65e5\u5148\u68c0\u67e5 8001\u30018002\u30018003 \u4e09\u4e2a\u4e3b\u529b\u4ed3\u5e93\u7684\u9500\u552e\u548c\u6bdb\u5229\u5dee\u5f02\u3002"),
            U(r"\u628a 7 \u6708\u76ee\u6807\u62c6\u5230\u6bcf\u5e97\u6bcf\u65e5\uff0c\u6bcf\u5929\u8ffd\u8e2a\u5dee\u989d\u3002"),
            U(r"\u5e93\u5b58\u5206\u6790\u4e0b\u4e00\u6b65\u8981\u52a0\u5165\u6ede\u9500\u5929\u6570\u548c\u5c3a\u7801\u7ed3\u6784\u3002"),
        ],
        "todos": [
            U(r"\u786e\u8ba4\u4ed3\u5e93\u4ee3\u7801 8001\u30018002\u30018003\u30018014 \u5bf9\u5e94\u7684\u95e8\u5e97\u540d\u79f0\u3002"),
            U(r"\u68c0\u67e5 SAP B1 \u51cc\u6668 2:00 \u81ea\u52a8\u540c\u6b65\u7ed3\u679c\u3002"),
            U(r"\u4e3a AI \u603b\u7ecf\u7406\u63a5\u5165\u6570\u636e\u67e5\u8be2\u5de5\u5177\u3002"),
        ],
    }
    try:
        if os.path.exists(SAP_SUMMARY_FILE):
            with open(SAP_SUMMARY_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            return {**fallback, **data}
    except Exception:
        pass
    return fallback


def get_secret():
    os.makedirs(APP_DIR, exist_ok=True)
    if not os.path.exists(SECRET_FILE):
        Path(SECRET_FILE).write_text(secrets.token_urlsafe(48), encoding="utf-8")
        os.chmod(SECRET_FILE, 0o600)
    return Path(SECRET_FILE).read_text(encoding="utf-8").strip().encode()


SECRET = get_secret()


def db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn


def hp(password, salt=None, iterations=200000):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), iterations)
    return salt + "$" + base64.b64encode(digest).decode()


def cp(password, stored):
    try:
        salt = stored.split("$", 1)[0]
        return hmac.compare_digest(hp(password, salt, 200000), stored) or hmac.compare_digest(hp(password, salt, 160000), stored)
    except Exception:
        return False


def needs_password_upgrade(password, stored):
    try:
        salt = stored.split("$", 1)[0]
        return hmac.compare_digest(hp(password, salt, 160000), stored) and not hmac.compare_digest(hp(password, salt, 200000), stored)
    except Exception:
        return False


def sign(value):
    digest = hmac.new(SECRET, value.encode(), hashlib.sha256).hexdigest()
    return value + "|" + digest


def unsign(value):
    if not value or "|" not in value:
        return None
    raw, digest = value.rsplit("|", 1)
    good = hmac.new(SECRET, raw.encode(), hashlib.sha256).hexdigest()
    return raw if hmac.compare_digest(digest, good) else None


def ensure_column(conn, table, column, ddl):
    cols = {r["name"] for r in conn.execute(f"pragma table_info({table})").fetchall()}
    if column not in cols:
        conn.execute(f"alter table {table} add column {ddl}")


def init():
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    with db() as conn:
        conn.execute(
            """
create table if not exists users(
 id integer primary key autoincrement,
 email text unique not null,
 name text not null,
 phone text,
 store text,
 role text not null default 'employee',
 status text not null default 'pending',
 password_hash text not null,
 created_at integer not null,
 last_login integer
)
"""
        )
        ensure_column(conn, "users", "failed_attempts", "failed_attempts integer not null default 0")
        ensure_column(conn, "users", "locked_until", "locked_until integer not null default 0")
        ensure_column(conn, "users", "updated_at", "updated_at integer")
        ensure_column(conn, "users", "reset_required", "reset_required integer not null default 0")
        for old, new in OLD_ROLE_MAP.items():
            conn.execute("update users set role=? where role=?", (new, old))
        conn.execute(
            """
create table if not exists records(
 id integer primary key autoincrement,
 module text not null,
 title text not null,
 status text not null default 'active',
 tags text,
 summary text,
 data_json text,
 created_by integer,
 created_at integer not null,
 updated_at integer not null
)
"""
        )
        conn.execute("create index if not exists idx_records_module on records(module)")
        conn.execute(
            """
create table if not exists knowledge_items(
 id integer primary key autoincrement,
 title text not null,
 category text,
 tags text,
 body text,
 ai_summary text,
 source_type text,
 source_ref text,
 approved integer not null default 1,
 created_by integer,
 created_at integer not null,
 updated_at integer not null
)
"""
        )
        conn.execute("create index if not exists idx_knowledge_category on knowledge_items(category)")
        conn.execute(
            """
create table if not exists uploaded_files(
 id integer primary key autoincrement,
 original_name text not null,
 saved_name text not null,
 path text not null,
 mime text,
 size integer,
 category text,
 description text,
 public integer not null default 1,
 need_summary integer not null default 1,
 need_sales_script integer not null default 0,
 extracted_text text,
 knowledge_id integer,
 created_by integer,
 created_at integer not null
)
"""
        )
        conn.execute(
            """
create table if not exists activity_log(
 id integer primary key autoincrement,
 user_id integer,
 action text not null,
 target_type text,
 target_id integer,
 detail text,
 created_at integer not null
)
"""
        )
        conn.execute(
            """
create table if not exists timeline_events(
 id integer primary key autoincrement,
 target_type text not null,
 target_id integer not null,
 title text not null,
 body text,
 created_by integer,
 created_at integer not null
)
"""
        )
        conn.execute(
            """
create table if not exists relations(
 id integer primary key autoincrement,
 from_type text not null,
 from_id integer not null,
 to_type text not null,
 to_id integer not null,
 relation_type text,
 created_by integer,
 created_at integer not null
)
"""
        )
        admin_email = os.environ.get("PORTAL_ADMIN_EMAIL", "vafox@126.com").strip().lower()
        existing_admin = conn.execute("select id from users where role='admin' limit 1").fetchone()
        if not existing_admin:
            initial_password = os.environ.get("PORTAL_INITIAL_ADMIN_PASSWORD")
            if initial_password:
                conn.execute(
                    "insert into users(email,name,phone,store,role,status,password_hash,created_at) values(?,?,?,?,?,?,?,?)",
                    (admin_email, U(r"\u7ba1\u7406\u5458"), "", U(r"\u603b\u90e8"), "admin", "approved", hp(initial_password), int(time.time())),
                )
        conn.commit()


init()


def layout(title, body, user=None, msg="", wide=False):
    nav = ""
    if user:
        nav = (
            '<div class="topbar"><div><strong>{}</strong><small>{} · {}</small></div>'
            '<div><a href="/change-password">{}</a><a href="/logout">{}</a></div></div>'
        ).format(esc(user["name"]), esc(ROLES.get(user["role"], user["role"])), esc(user["store"]), T["change_password"], T["logout"])
    alert = f'<div class="alert">{esc(msg)}</div>' if msg else ""
    max_width = "1180px" if wide else "980px"
    return f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{esc(title)}</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;font-family:-apple-system,BlinkMacSystemFont,"Segoe UI","Microsoft YaHei",Arial,sans-serif;background:#f6f3ed;color:#171717}}
main{{width:min({max_width},calc(100% - 24px));margin:0 auto;padding:18px 0 44px}}
h1{{font-size:30px;margin:8px 0 6px}}h2{{font-size:20px;margin:0 0 12px}}p,.lead{{line-height:1.65}}.lead{{font-size:16px;color:#555;margin:0 0 16px}}
.topbar{{display:flex;justify-content:space-between;gap:12px;align-items:center;background:#fff;border:1px solid #ddd7cc;border-radius:8px;padding:12px 14px;margin-bottom:18px}}
.topbar small{{display:block;color:#666;margin-top:3px}}.topbar a{{margin-left:12px;color:#1849a9;text-decoration:none;font-weight:700}}
.panel,.card{{background:#fff;border:1px solid #ddd7cc;border-radius:8px;box-shadow:0 8px 22px rgba(0,0,0,.05)}}.panel{{padding:18px;margin:14px 0}}.form{{max-width:520px}}
.grid{{display:grid;grid-template-columns:repeat(3,1fr);gap:14px}}.card{{padding:18px;min-height:154px;display:flex;flex-direction:column;justify-content:space-between}}
.metrics{{display:grid;grid-template-columns:repeat(5,1fr);gap:10px;margin:12px 0}}.metric{{background:#fff;border:1px solid #ddd7cc;border-radius:8px;padding:14px;min-height:92px}}.metric strong{{display:block;font-size:22px;margin-top:7px}}.metric span{{font-size:13px;color:#666}}
.split{{display:grid;grid-template-columns:1fr 1fr;gap:14px}}.list{{margin:0;padding-left:20px;line-height:1.85}}.pill{{display:inline-block;border:1px solid #ddd7cc;border-radius:999px;padding:7px 10px;margin:3px 5px 3px 0;background:#fff;font-weight:700;color:#333;text-decoration:none}}
.store-row{{display:grid;grid-template-columns:1.1fr 1fr 1fr;gap:8px;border-top:1px solid #eee;padding:10px 0}}.store-row:first-child{{border-top:0}}
.card h2{{margin-bottom:4px}}.card p{{color:#555;margin:0 0 18px}}.disabled{{opacity:.55}}
label{{display:block;font-weight:800;margin:12px 0 7px}}input,select,textarea{{width:100%;padding:14px;border:1px solid #cfc8bb;border-radius:8px;font-size:16px;background:#fff}}textarea{{min-height:120px;font-family:inherit}}
button,.btn{{display:inline-block;border:0;border-radius:8px;background:#1849a9;color:#fff;text-decoration:none;font-weight:800;padding:13px 16px;cursor:pointer;font-size:16px;text-align:center}}
.btn.full{{width:100%}}.red{{background:#ad1f15}}.green{{background:#18704c}}.dark{{background:#222}}.gray{{background:#777}}.orange{{background:#b45f06}}
.alert{{padding:12px;background:#fff7d6;border:1px solid #ecd27a;border-radius:8px;margin:12px 0}}table{{width:100%;border-collapse:collapse}}th,td{{border-bottom:1px solid #eee;padding:10px;text-align:left;vertical-align:top}}th{{white-space:nowrap}}.inline{{display:flex;gap:8px;align-items:center;flex-wrap:wrap}}.inline form{{display:inline}}.small{{font-size:13px;color:#666}}
@media(max-width:820px){{main{{width:calc(100% - 18px);padding-top:10px}}h1{{font-size:26px}}.grid,.metrics,.split{{grid-template-columns:1fr;gap:12px}}.store-row{{grid-template-columns:1fr}}.card{{min-height:132px}}.btn,button{{width:100%;padding:15px}}.topbar{{align-items:flex-start;flex-direction:column}}.topbar a{{margin:0 12px 0 0}}table,tbody,tr,td,th{{display:block}}thead{{display:none}}tr{{border:1px solid #eee;border-radius:8px;margin:10px 0;padding:8px;background:#fff}}td{{border:0;padding:7px}}.inline{{display:block}}.inline form{{display:block;margin-top:8px}}}}
</style></head><body><main>{nav}<section><h1>{esc(title)}</h1><p class="lead">{T['subtitle']}</p></section>{alert}{body}</main></body></html>"""


class App(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        print(fmt % args)

    def out(self, html_text, code=200):
        body = html_text.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Frame-Options", "SAMEORIGIN")
        self.end_headers()
        self.wfile.write(body)

    def json_out(self, data, code=200):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def file_out(self, name, content, content_type="text/csv; charset=utf-8"):
        body = content.encode("utf-8-sig")
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{name}"')
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def redir(self, path, cookie=None):
        self.send_response(302)
        self.send_header("Location", path)
        if cookie:
            self.send_header("Set-Cookie", cookie)
        self.end_headers()

    def form(self):
        size = int(self.headers.get("Content-Length", "0") or 0)
        return {k: v[0] for k, v in parse_qs(self.rfile.read(size).decode("utf-8")).items()}

    def multipart(self):
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST"})
        return form

    def current_user(self):
        jar = cookies.SimpleCookie(self.headers.get("Cookie", ""))
        token = jar.get("fp_session")
        uid = unsign(token.value) if token else None
        if not uid:
            return None
        with db() as conn:
            return conn.execute("select * from users where id=? and status='approved'", (uid,)).fetchone()

    def require_admin(self):
        user = self.current_user()
        if not user or user["role"] != "admin":
            self.redir("/login")
            return None
        return user

    def do_GET(self):
        path = urlparse(self.path).path
        user = self.current_user()
        if path == "/logout":
            return self.redir("/login", "fp_session=; Path=/; Max-Age=0; HttpOnly; Secure; SameSite=Lax")
        if path == "/login":
            return self.login()
        if path == "/register":
            return self.register()
        if path == "/change-password":
            return self.change_password(user)
        if path == "/admin":
            return self.admin()
        if path == "/ai-ceo":
            return self.ai_ceo(user)
        if path == "/ai-store-manager":
            return self.ai_store_manager(user)
        if path == "/ai-query":
            return self.ai_assistant(user)
        if path == "/agents":
            return self.agents(user)
        if path == "/sap-sync":
            return self.sap_sync(user)
        if path == "/documents":
            return self.document_center(user)
        if path == "/workflow":
            return self.workflow_center(user)
        if path == "/business-overview":
            return self.business_overview(user)
        if path == "/content-center":
            return self.module_page(user, "/content")
        if path == "/knowledge":
            return self.knowledge(user)
        if path == "/inventory":
            return self.inventory(user)
        if path == "/records/new":
            return self.record_form(user)
        if path == "/records/view":
            return self.record_view(user)
        if path == "/records/edit":
            return self.record_form(user, edit=True)
        if path == "/records/export":
            return self.records_export(user)
        if path in MODULES:
            return self.module_page(user, path)
        if path == "/upload":
            return self.upload(user)
        if path == "/web-search":
            return self.web_search(user)
        if path == "/ai-assistant":
            return self.ai_assistant(user)
        if path == "/knowledge/view":
            return self.knowledge_view(user)
        if path == "/knowledge/new":
            return self.knowledge_form(user)
        if path == "/api/dashboard/summary":
            return self.json_out(load_summary())
        if path.startswith("/api/sap/"):
            return self.sap_api_placeholder(user, path)
        if path == "/":
            return self.dashboard(user) if user else self.login()
        if path == "/daily":
            return self.redir("/wiki/firefox-hq/daily")
        return self.redir("/")

    def do_POST(self):
        path = urlparse(self.path).path
        if path.startswith("/api/"):
            return self.json_out({"ok": False, "message": U(r"\u63a5\u53e3\u5df2\u9884\u7559\uff0c\u4e0b\u4e00\u6b65\u63a5\u5165 AI \u548c\u77e5\u8bc6\u5e93\u6570\u636e\u3002")}, code=501)
        if path == "/login":
            return self.login_post()
        if path == "/register":
            return self.register_post()
        if path == "/change-password":
            return self.change_password_post()
        if path == "/records/save":
            return self.record_save()
        if path == "/records/delete":
            return self.record_delete()
        if path == "/records/import":
            return self.records_import()
        if path == "/upload":
            return self.upload_post()
        if path == "/web-search":
            return self.web_search_post()
        if path in ("/ai-assistant", "/ai-query"):
            return self.ai_assistant_post()
        if path == "/knowledge/save":
            return self.knowledge_save()
        if path == "/admin/update":
            return self.admin_update()
        if path == "/admin/reset-password":
            return self.admin_reset_password()
        return self.redir("/")

    def login(self, msg=""):
        body = f"""
<div class="panel form">
  <form method="post" action="/login">
    <label>{T['email']}</label><input name="email" type="email" autocomplete="username" required>
    <label>{T['password']}</label><input name="password" type="password" autocomplete="current-password" required>
    <p><button>{T['login']}</button></p>
  </form>
  <p><a href="/register">{T['register']}</a></p>
</div>"""
        self.out(layout(T["brand"], body, msg=msg))

    def login_post(self):
        form = self.form()
        email = form.get("email", "").strip().lower()
        password = form.get("password", "")
        now = int(time.time())
        with db() as conn:
            user = conn.execute("select * from users where email=?", (email,)).fetchone()
            if user and user["locked_until"] and user["locked_until"] > now:
                return self.login(T["locked"])
            ok = bool(user and user["status"] == "approved" and cp(password, user["password_hash"]))
            if ok:
                new_hash = hp(password) if needs_password_upgrade(password, user["password_hash"]) else user["password_hash"]
                conn.execute(
                    "update users set password_hash=?, failed_attempts=0, locked_until=0, last_login=?, updated_at=? where id=?",
                    (new_hash, now, now, user["id"]),
                )
                return self.redir("/", "fp_session={}; Path=/; HttpOnly; Secure; SameSite=Lax".format(sign(str(user["id"]))))
            if user:
                attempts = int(user["failed_attempts"] or 0) + 1
                locked_until = now + LOCK_SECONDS if attempts >= LOCK_LIMIT else 0
                conn.execute(
                    "update users set failed_attempts=?, locked_until=?, updated_at=? where id=?",
                    (attempts, locked_until, now, user["id"]),
                )
                if locked_until:
                    return self.login(T["locked"])
        return self.login(T["bad"])

    def register(self, msg=""):
        body = f"""
<div class="panel form">
  <form method="post" action="/register">
    <label>{T['name']}</label><input name="name" required>
    <label>{T['phone']}</label><input name="phone">
    <label>{T['store']}</label><input name="store">
    <label>{T['email']}</label><input name="email" type="email" autocomplete="username" required>
    <label>{T['password']}</label><input name="password" type="password" autocomplete="new-password" required>
    <p><button>{T['register']}</button></p>
  </form>
  <p><a href="/login">{T['login']}</a></p>
</div>"""
        self.out(layout(T["register"], body, msg=msg))

    def register_post(self):
        form = self.form()
        try:
            with db() as conn:
                conn.execute(
                    "insert into users(email,name,phone,store,role,status,password_hash,created_at,updated_at) values(?,?,?,?,?,?,?,?,?)",
                    (
                        form.get("email", "").strip().lower(),
                        form.get("name", "").strip(),
                        form.get("phone", "").strip(),
                        form.get("store", "").strip(),
                        "employee",
                        "pending",
                        hp(form.get("password", "")),
                        int(time.time()),
                        int(time.time()),
                    ),
                )
            return self.login(T["pending"])
        except sqlite3.IntegrityError:
            return self.register(T["duplicate"])

    def card(self, title, text, href, cls="btn", allow=True):
        if allow:
            action = f'<a class="{cls} full" href="{esc(href)}">{esc(title)}</a>'
            extra = ""
        else:
            action = f'<span class="btn gray full">{T["no_permission"]}</span>'
            extra = " disabled"
        return f'<div class="card{extra}"><div><h2>{esc(title)}</h2><p>{esc(text)}</p></div>{action}</div>'

    def metric(self, label, value, note=""):
        return '<div class="metric"><span>{}</span><strong>{}</strong><span>{}</span></div>'.format(esc(label), esc(value), esc(note))

    def bullets(self, items):
        clean = [item for item in items if item]
        return '<ul class="list">' + "".join("<li>{}</li>".format(esc(item)) for item in clean) + "</ul>"

    def require_login(self, user):
        if not user:
            self.redir("/login")
            return None
        return user

    def can_open(self, user, roles):
        return bool(user and user["role"] in roles)

    def log_action(self, user, action, target_type="", target_id=None, detail=""):
        try:
            with db() as conn:
                conn.execute(
                    "insert into activity_log(user_id,action,target_type,target_id,detail,created_at) values(?,?,?,?,?,?)",
                    (user["id"] if user else None, action, target_type, target_id, detail, ts()),
                )
        except Exception:
            pass

    def record_allowed(self, user, module):
        data = MODULES.get("/" + module_key(module))
        return bool(data and self.can_open(user, data[2]))

    def placeholder(self, user, title, text):
        user = self.require_login(user)
        if not user:
            return
        api_list = [
            "POST /api/upload",
            "POST /api/knowledge/create",
            "GET /api/knowledge/list",
            "GET /api/knowledge/:id",
            "POST /api/search/web",
            "POST /api/search/save",
            "POST /api/ai/chat",
            "POST /api/ai/summarize",
            "POST /api/ai/classify",
            "GET /api/dashboard/summary",
        ]
        body = f"""
<div class="panel">
  <h2>{esc(title)}</h2>
  <p>{esc(text)}</p>
  <h2>{U(r'\u5df2\u9884\u7559\u63a5\u53e3')}</h2>
  {self.bullets(api_list)}
  <p><a class="btn" href="/">{U(r'\u8fd4\u56de\u9996\u9875')}</a></p>
</div>"""
        self.out(layout(title, body, user=user))

    def module_page(self, user, path):
        user = self.require_login(user)
        if not user:
            return
        module = module_key(path)
        title, text, roles = MODULES[path]
        if not self.can_open(user, roles):
            return self.dashboard(user)
        q = parse_qs(urlparse(self.path).query).get("q", [""])[0].strip()
        sql = "select * from records where module=? and status!='deleted'"
        params = [module]
        if q:
            sql += " and (title like ? or tags like ? or summary like ? or data_json like ?)"
            like = "%" + q + "%"
            params += [like, like, like, like]
        sql += " order by updated_at desc limit 100"
        with db() as conn:
            rows = conn.execute(sql, params).fetchall()
        abilities = [
            U(r"\u65b0\u5efa / \u7f16\u8f91 / \u67e5\u770b\u8be6\u60c5"),
            U(r"\u641c\u7d22 / \u7b5b\u9009 / \u6807\u7b7e\u7ba1\u7406"),
            U(r"\u4e0a\u4f20\u56fe\u7247\u3001PDF\u3001Word\u3001Excel\u3001PPT\u3001TXT\u3001Markdown"),
            U(r"\u6587\u4ef6\u5939 / \u81ea\u52a8\u5f52\u6863 / \u65f6\u95f4\u8f74"),
            U(r"AI \u603b\u7ed3 / AI \u5efa\u8bae / AI \u67e5\u8be2"),
            U(r"\u64cd\u4f5c\u65e5\u5fd7 / \u6743\u9650\u63a7\u5236 / Excel \u5bfc\u5165\u5bfc\u51fa"),
        ]
        items = ""
        for row in rows:
            items += "<tr><td>{}</td><td>{}</td><td>{}</td><td><a href=\"/records/view?id={}\">{}</a> <a href=\"/records/edit?id={}\">{}</a></td></tr>".format(
                esc(row["title"]),
                esc(row["tags"]),
                esc(dt(row["updated_at"])),
                row["id"],
                U(r"\u8be6\u60c5"),
                row["id"],
                U(r"\u7f16\u8f91"),
            )
        if not items:
            items = '<tr><td colspan="4" class="small">{}</td></tr>'.format(U(r"\u6682\u65e0\u6863\u6848\uff0c\u53ef\u4ee5\u5148\u65b0\u5efa\u4e00\u6761\u3002"))
        body = f"""
<div class="panel">
  <h2>{esc(title)}</h2>
  <p>{esc(text)}</p>
  <div class="inline">
    <a class="btn" href="/records/new?module={esc(module)}">{U(r'\u65b0\u5efa')}</a>
    <a class="btn green" href="/records/export?module={esc(module)}">{U(r'Excel \u5bfc\u51fa')}</a>
    <a class="btn orange" href="/upload">{U(r'\u4e0a\u4f20\u9644\u4ef6')}</a>
  </div>
</div>
<div class="panel">
  <form method="get" action="/{esc(module)}">
    <label>{U(r'\u641c\u7d22')}</label><input name="q" value="{esc(q)}" placeholder="{U(r'\u8f93\u5165\u540d\u79f0\u3001\u6807\u7b7e\u6216\u5185\u5bb9')}">
    <p><button>{U(r'\u641c\u7d22')}</button></p>
  </form>
  <table><thead><tr><th>{U(r'\u540d\u79f0')}</th><th>{U(r'\u6807\u7b7e')}</th><th>{U(r'\u66f4\u65b0')}</th><th>{T['action']}</th></tr></thead><tbody>{items}</tbody></table>
</div>
<div class="panel">
  <h2>{U(r'\u901a\u7528\u80fd\u529b')}</h2>
  {self.bullets(abilities)}
  <form method="post" action="/records/import">
    <input type="hidden" name="module" value="{esc(module)}">
    <label>{U(r'Excel/CSV \u6279\u91cf\u5bfc\u5165')}</label>
    <textarea name="csv_text" rows="5" placeholder="{U(r'\u7c98\u8d34 Excel \u590d\u5236\u7684\u8868\u683c\uff1a\u7b2c\u4e00\u5217\u4e3a\u540d\u79f0\uff0c\u7b2c\u4e8c\u5217\u4e3a\u6807\u7b7e\uff0c\u7b2c\u4e09\u5217\u4e3a\u5907\u6ce8')}"></textarea>
    <p><button>{U(r'\u6279\u91cf\u5bfc\u5165')}</button></p>
  </form>
</div>"""
        self.out(layout(title, body, user=user))

    def record_form(self, user, edit=False):
        user = self.require_login(user)
        if not user:
            return
        query = parse_qs(urlparse(self.path).query)
        row = None
        module = module_key(query.get("module", [""])[0])
        if edit:
            rid = query.get("id", [""])[0]
            with db() as conn:
                row = conn.execute("select * from records where id=? and status!='deleted'", (rid,)).fetchone()
            if not row:
                return self.redir("/")
            module = row["module"]
        if not self.record_allowed(user, module):
            return self.dashboard(user)
        title = row["title"] if row else ""
        tags = row["tags"] if row else ""
        summary = row["summary"] if row else ""
        data = json.loads(row["data_json"] or "{}") if row else {}
        body = f"""
<div class="panel form">
  <form method="post" action="/records/save">
    <input type="hidden" name="id" value="{esc(row['id'] if row else '')}">
    <input type="hidden" name="module" value="{esc(module)}">
    <label>{U(r'\u540d\u79f0')}</label><input name="title" value="{esc(title)}" required>
    <label>{U(r'\u6807\u7b7e')}</label><input name="tags" value="{esc(tags)}" placeholder="{U(r'\u591a\u4e2a\u6807\u7b7e\u7528\u9017\u53f7\u5206\u9694')}">
    <label>{U(r'\u6458\u8981')}</label><textarea name="summary">{esc(summary)}</textarea>
    <label>{U(r'\u6838\u5fc3\u5b57\u6bb51')}</label><input name="field1" value="{esc(data.get('field1',''))}">
    <label>{U(r'\u6838\u5fc3\u5b57\u6bb52')}</label><input name="field2" value="{esc(data.get('field2',''))}">
    <label>{U(r'\u8be6\u7ec6\u5185\u5bb9')}</label><textarea name="body">{esc(data.get('body',''))}</textarea>
    <p><button>{T['save']}</button></p>
  </form>
</div>"""
        self.out(layout((U(r"\u7f16\u8f91") if edit else U(r"\u65b0\u5efa")) + " - " + module_title(module), body, user=user))

    def record_save(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.form()
        module = module_key(form.get("module"))
        if not self.record_allowed(user, module):
            return self.dashboard(user)
        rid = form.get("id", "").strip()
        title = form.get("title", "").strip()
        data = {"field1": form.get("field1", ""), "field2": form.get("field2", ""), "body": form.get("body", "")}
        now = ts()
        with db() as conn:
            if rid:
                conn.execute(
                    "update records set title=?, tags=?, summary=?, data_json=?, updated_at=? where id=?",
                    (title, form.get("tags", ""), form.get("summary", ""), json.dumps(data, ensure_ascii=False), now, rid),
                )
                target_id = int(rid)
                action = "record_update"
            else:
                cur = conn.execute(
                    "insert into records(module,title,tags,summary,data_json,created_by,created_at,updated_at) values(?,?,?,?,?,?,?,?)",
                    (module, title, form.get("tags", ""), form.get("summary", ""), json.dumps(data, ensure_ascii=False), user["id"], now, now),
                )
                target_id = cur.lastrowid
                action = "record_create"
            conn.execute(
                "insert into timeline_events(target_type,target_id,title,body,created_by,created_at) values(?,?,?,?,?,?)",
                ("record", target_id, U(r"\u6863\u6848\u5df2\u4fdd\u5b58"), title, user["id"], now),
            )
        self.log_action(user, action, "record", target_id, title)
        return self.redir(f"/records/view?id={target_id}")

    def record_view(self, user):
        user = self.require_login(user)
        if not user:
            return
        rid = parse_qs(urlparse(self.path).query).get("id", [""])[0]
        with db() as conn:
            row = conn.execute("select * from records where id=? and status!='deleted'", (rid,)).fetchone()
            events = conn.execute("select * from timeline_events where target_type='record' and target_id=? order by created_at desc limit 20", (rid,)).fetchall()
            logs = conn.execute("select * from activity_log where target_type='record' and target_id=? order by created_at desc limit 20", (rid,)).fetchall()
        if not row or not self.record_allowed(user, row["module"]):
            return self.redir("/")
        data = json.loads(row["data_json"] or "{}")
        event_html = self.bullets([dt(e["created_at"]) + " " + e["title"] for e in events]) if events else "<p class='small'>暂无时间轴。</p>"
        log_html = self.bullets([dt(l["created_at"]) + " " + l["action"] for l in logs]) if logs else "<p class='small'>暂无操作日志。</p>"
        body = f"""
<div class="panel">
  <h2>{esc(row['title'])}</h2>
  <p class="small">{module_title(row['module'])} ｜ {U(r'\u6807\u7b7e')}：{esc(row['tags'])} ｜ {U(r'\u66f4\u65b0')}：{esc(dt(row['updated_at']))}</p>
  <p>{esc(row['summary'])}</p>
  <div class="split">
    <div><h2>{U(r'\u8be6\u7ec6\u5185\u5bb9')}</h2><p>{esc(data.get('body',''))}</p></div>
    <div><h2>{U(r'AI \u5efa\u8bae')}</h2>{self.bullets([U(r'\u5b8c\u5584\u56fe\u7247\u3001\u9644\u4ef6\u548c\u5173\u8054\u5173\u7cfb\u3002'), U(r'\u8865\u5145\u9500\u552e\u6570\u636e\u548c\u5386\u53f2\u8bb0\u5f55\uff0c\u4fbf\u4e8e\u540e\u7eed AI \u5206\u6790\u3002')])}</div>
  </div>
  <div class="inline"><a class="btn" href="/records/edit?id={row['id']}">{U(r'\u7f16\u8f91')}</a><a class="btn gray" href="/{esc(row['module'])}">{U(r'\u8fd4\u56de')}</a><form method="post" action="/records/delete"><input type="hidden" name="id" value="{row['id']}"><button class="red">{U(r'\u5220\u9664')}</button></form></div>
</div>
<div class="split">
  <div class="panel"><h2>{U(r'\u65f6\u95f4\u8f74')}</h2>{event_html}</div>
  <div class="panel"><h2>{U(r'\u64cd\u4f5c\u65e5\u5fd7')}</h2>{log_html}</div>
</div>"""
        self.out(layout(row["title"], body, user=user, wide=True))

    def record_delete(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.form()
        rid = form.get("id")
        with db() as conn:
            row = conn.execute("select * from records where id=?", (rid,)).fetchone()
            if row and self.record_allowed(user, row["module"]):
                conn.execute("update records set status='deleted', updated_at=? where id=?", (ts(), rid))
                self.log_action(user, "record_delete", "record", rid, row["title"])
                return self.redir("/" + row["module"])
        return self.redir("/")

    def records_export(self, user):
        user = self.require_login(user)
        if not user:
            return
        module = module_key(parse_qs(urlparse(self.path).query).get("module", [""])[0])
        if not self.record_allowed(user, module):
            return self.dashboard(user)
        with db() as conn:
            rows = conn.execute("select * from records where module=? and status!='deleted' order by updated_at desc", (module,)).fetchall()
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([U(r"\u540d\u79f0"), U(r"\u6807\u7b7e"), U(r"\u6458\u8981"), U(r"\u66f4\u65b0\u65f6\u95f4")])
        for row in rows:
            writer.writerow([row["title"], row["tags"], row["summary"], dt(row["updated_at"])])
        return self.file_out(module + "_export.csv", buf.getvalue())

    def records_import(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.form()
        module = module_key(form.get("module"))
        if not self.record_allowed(user, module):
            return self.dashboard(user)
        text = form.get("csv_text", "")
        count = 0
        now = ts()
        with db() as conn:
            for row in csv.reader(io.StringIO(text), delimiter="\t"):
                if not row:
                    continue
                if len(row) == 1 and "," in row[0]:
                    row = next(csv.reader([row[0]]))
                title = (row[0] if len(row) > 0 else "").strip()
                if not title or title in (U(r"\u540d\u79f0"), "title"):
                    continue
                tags = row[1].strip() if len(row) > 1 else ""
                summary = row[2].strip() if len(row) > 2 else ""
                conn.execute(
                    "insert into records(module,title,tags,summary,data_json,created_by,created_at,updated_at) values(?,?,?,?,?,?,?,?)",
                    (module, title, tags, summary, "{}", user["id"], now, now),
                )
                count += 1
        self.log_action(user, "records_import", "record", None, f"{module}:{count}")
        return self.redir("/" + module)

    def upload(self, user):
        user = self.require_login(user)
        if not user:
            return
        module_opts = "".join('<option value="{}">{}</option>'.format(module_key(k), esc(v[0])) for k, v in MODULES.items())
        body = f"""
<div class="panel form">
  <h2>{U(r'\u6587\u4ef6\u4e0a\u4f20\u4e0e\u81ea\u52a8\u5f52\u6863')}</h2>
  <form method="post" action="/upload" enctype="multipart/form-data">
    <label>{U(r'\u6587\u4ef6')}</label><input name="file" type="file" required>
    <label>{U(r'\u5f52\u5c5e\u6a21\u5757')}</label><select name="module"><option value="knowledge">{U(r'\u77e5\u8bc6\u4e2d\u5fc3')}</option>{module_opts}</select>
    <label>{U(r'\u6240\u5c5e\u5206\u7c7b')}</label><input name="category" placeholder="01_公司制度 / 02_产品资料 / 07_库存采购">
    <label>{U(r'\u8865\u5145\u8bf4\u660e')}</label><textarea name="description"></textarea>
    <p><button>{U(r'\u4e0a\u4f20\u5e76\u81ea\u52a8\u5165\u5e93')}</button></p>
  </form>
</div>"""
        self.out(layout(U(r"\u6587\u4ef6\u4e0a\u4f20"), body, user=user))

    def upload_post(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.multipart()
        item = form["file"] if "file" in form else None
        if not item or not getattr(item, "filename", ""):
            return self.upload(user)
        original = Path(item.filename).name
        module = module_key(form.getfirst("module", "knowledge"))
        category = form.getfirst("category", "").strip()
        description = form.getfirst("description", "").strip()
        folder = Path(UPLOAD_DIR) / module
        folder.mkdir(parents=True, exist_ok=True)
        saved = uuid.uuid4().hex + Path(original).suffix.lower()
        path = folder / saved
        data = item.file.read()
        path.write_bytes(data)
        extracted = extract_file_text(str(path), original)
        combined = "\n".join([original, description, extracted])
        category = category or classify_text(combined)
        summary = summarize_text(combined)
        tags = extract_tags(combined)
        now = ts()
        with db() as conn:
            kcur = conn.execute(
                "insert into knowledge_items(title,category,tags,body,ai_summary,source_type,source_ref,created_by,created_at,updated_at) values(?,?,?,?,?,?,?,?,?,?)",
                (original, category, tags, extracted or description, summary, U(r"\u6587\u4ef6\u4e0a\u4f20"), original, user["id"], now, now),
            )
            kid = kcur.lastrowid
            conn.execute(
                "insert into uploaded_files(original_name,saved_name,path,mime,size,category,description,extracted_text,knowledge_id,created_by,created_at) values(?,?,?,?,?,?,?,?,?,?,?)",
                (original, saved, str(path), mimetypes.guess_type(original)[0], len(data), category, description, extracted, kid, user["id"], now),
            )
        self.log_action(user, "file_upload", "knowledge", kid, original)
        return self.redir(f"/knowledge/view?id={kid}")

    def knowledge_form(self, user):
        user = self.require_login(user)
        if not user:
            return
        body = f"""
<div class="panel form">
  <h2>{U(r'\u65b0\u5efa\u77e5\u8bc6')}</h2>
  <form method="post" action="/knowledge/save">
    <label>{U(r'\u6807\u9898')}</label><input name="title" required>
    <label>{U(r'\u5206\u7c7b')}</label><input name="category" placeholder="01_公司制度">
    <label>{U(r'\u6807\u7b7e')}</label><input name="tags">
    <label>{U(r'\u6b63\u6587')}</label><textarea name="body" required></textarea>
    <p><button>{T['save']}</button></p>
  </form>
</div>"""
        self.out(layout(U(r"\u65b0\u5efa\u77e5\u8bc6"), body, user=user))

    def knowledge_save(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.form()
        body = form.get("body", "")
        category = form.get("category", "").strip() or classify_text(body)
        now = ts()
        with db() as conn:
            cur = conn.execute(
                "insert into knowledge_items(title,category,tags,body,ai_summary,source_type,source_ref,created_by,created_at,updated_at) values(?,?,?,?,?,?,?,?,?,?)",
                (form.get("title", "").strip(), category, form.get("tags", "") or extract_tags(body), body, summarize_text(body), U(r"\u624b\u52a8\u5f55\u5165"), "", user["id"], now, now),
            )
            kid = cur.lastrowid
        self.log_action(user, "knowledge_create", "knowledge", kid, form.get("title", ""))
        return self.redir(f"/knowledge/view?id={kid}")

    def knowledge_view(self, user):
        user = self.require_login(user)
        if not user:
            return
        kid = parse_qs(urlparse(self.path).query).get("id", [""])[0]
        with db() as conn:
            row = conn.execute("select * from knowledge_items where id=?", (kid,)).fetchone()
        if not row:
            return self.redir("/knowledge")
        body = f"""
<div class="panel">
  <h2>{esc(row['title'])}</h2>
  <p class="small">{esc(row['category'])} ｜ {esc(row['tags'])} ｜ {esc(row['source_type'])}：{esc(row['source_ref'])}</p>
  <h2>{U(r'AI \u6458\u8981')}</h2><p>{esc(row['ai_summary'])}</p>
  <h2>{U(r'\u6b63\u6587')}</h2><p>{esc(row['body'])}</p>
  <p><a class="btn" href="/ai-query">{U(r'\u7528 AI \u67e5\u8be2')}</a> <a class="btn gray" href="/knowledge">{U(r'\u8fd4\u56de')}</a></p>
</div>"""
        self.out(layout(row["title"], body, user=user, wide=True))

    def ai_assistant(self, user, answer="", question="", refs=None):
        user = self.require_login(user)
        if not user:
            return
        refs = refs or []
        ref_html = self.bullets(refs) if refs else "<p class='small'>暂无引用，提问后会显示来源。</p>"
        body = f"""
<div class="panel form">
  <h2>{U(r'AI \u667a\u80fd\u4f53\u67e5\u8be2')}</h2>
  <form method="post" action="/ai-query">
    <label>{U(r'\u95ee\u9898')}</label><textarea name="question" required>{esc(question)}</textarea>
    <label>{U(r'\u8303\u56f4')}</label><select name="scope"><option value="all">{U(r'\u5168\u516c\u53f8')}</option><option value="stores">{U(r'\u95e8\u5e97')}</option><option value="products">{U(r'\u4ea7\u54c1')}</option><option value="finance">{U(r'\u8d22\u52a1')}</option></select>
    <p><button>{U(r'\u67e5\u8be2')}</button></p>
  </form>
</div>
<div class="panel"><h2>{U(r'AI \u56de\u7b54')}</h2><p>{esc(answer or U(r'\u8bf7\u8f93\u5165\u95ee\u9898\uff0c\u7cfb\u7edf\u4f1a\u5148\u67e5\u5185\u90e8\u77e5\u8bc6\u5e93\u3002'))}</p><h2>{U(r'\u5f15\u7528\u6765\u6e90')}</h2>{ref_html}</div>"""
        self.out(layout(U(r"AI \u667a\u80fd\u4f53\u67e5\u8be2"), body, user=user, wide=True))

    def ai_assistant_post(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.form()
        question = form.get("question", "").strip()
        words = [w for w in re.split(r"\s+", question) if w][:5]
        with db() as conn:
            rows = []
            for w in words or [question]:
                rows += conn.execute("select * from knowledge_items where title like ? or body like ? or tags like ? order by updated_at desc limit 5", (f"%{w}%", f"%{w}%", f"%{w}%")).fetchall()
        seen, refs, snippets = set(), [], []
        for row in rows:
            if row["id"] in seen:
                continue
            seen.add(row["id"])
            refs.append(f"{row['title']} / {row['category']}")
            snippets.append(row["ai_summary"] or summarize_text(row["body"], 120))
        if snippets:
            answer = U(r"\u6839\u636e\u5185\u90e8\u77e5\u8bc6\u5e93\uff0c\u521d\u6b65\u7ed3\u8bba\uff1a") + " ".join(snippets[:3])
        else:
            answer = U(r"\u5185\u90e8\u77e5\u8bc6\u5e93\u6682\u672a\u547d\u4e2d\u8db3\u591f\u5185\u5bb9\uff0c\u5efa\u8bae\u5148\u4e0a\u4f20\u76f8\u5173\u6587\u6863\uff0c\u6216\u5230\u5916\u7f51\u641c\u7d22\u9875\u6293\u53d6\u8d44\u6599\u540e\u5165\u5e93\u3002")
        self.log_action(user, "ai_query", "knowledge", None, question)
        return self.ai_assistant(user, answer, question, refs)

    def web_search(self, user, msg=""):
        user = self.require_login(user)
        if not user:
            return
        body = f"""
<div class="panel form">
  <h2>{U(r'\u5916\u7f51\u8d44\u6599\u4fdd\u5b58\u5230\u77e5\u8bc6\u5e93')}</h2>
  <p class="small">{U(r'\u7b2c\u4e00\u7248\u652f\u6301\u7c98\u8d34\u7f51\u9875 URL \u6293\u53d6\u6b63\u6587\uff0c\u540e\u7eed\u518d\u63a5 Bing/SerpAPI/Tavily\u3002')}</p>
  <form method="post" action="/web-search">
    <label>URL</label><input name="url" placeholder="https://...">
    <label>{U(r'\u6807\u9898')}</label><input name="title">
    <p><button>{U(r'\u6293\u53d6\u5e76\u4fdd\u5b58')}</button></p>
  </form>
</div>"""
        self.out(layout(U(r"\u5916\u7f51\u641c\u7d22"), body, user=user, msg=msg))

    def web_search_post(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.form()
        url = form.get("url", "").strip()
        text = fetch_url_text(url)
        if not text:
            return self.web_search(user, U(r"\u6293\u53d6\u5931\u8d25\uff0c\u8bf7\u68c0\u67e5 URL\u3002"))
        title = form.get("title", "").strip() or summarize_text(text, 40)
        now = ts()
        with db() as conn:
            cur = conn.execute(
                "insert into knowledge_items(title,category,tags,body,ai_summary,source_type,source_ref,created_by,created_at,updated_at) values(?,?,?,?,?,?,?,?,?,?)",
                (title, classify_text(text), extract_tags(text), text, summarize_text(text), U(r"\u5916\u7f51\u641c\u7d22"), url, user["id"], now, now),
            )
            kid = cur.lastrowid
        self.log_action(user, "web_save", "knowledge", kid, url)
        return self.redir(f"/knowledge/view?id={kid}")

    def agents(self, user):
        user = self.require_login(user)
        if not user:
            return
        agents = [U(r"AI \u603b\u7ecf\u7406"), U(r"AI \u8d22\u52a1\u603b\u76d1"), U(r"AI \u91c7\u8d2d\u7ecf\u7406"), U(r"AI \u5e93\u5b58\u7ecf\u7406"), U(r"AI \u54c1\u724c\u7ecf\u7406"), U(r"AI \u95e8\u5e97\u7ecf\u7406"), U(r"AI \u4eba\u4e8b\u7ecf\u7406"), U(r"AI \u5185\u5bb9\u7ecf\u7406"), U(r"AI \u5ba2\u670d"), U(r"AI \u6cd5\u52a1"), U(r"AI \u6570\u636e\u5206\u6790\u5e08"), U(r"AI \u6295\u8d44\u52a9\u624b")]
        cards = "".join(self.card(a, U(r"\u57fa\u4e8e\u6743\u9650\u8303\u56f4\u67e5\u8be2\u6570\u636e\uff0c\u8f93\u51fa\u5efa\u8bae\u548c\u884c\u52a8\u6e05\u5355\u3002"), "/ai-query", "btn", True) for a in agents)
        self.out(layout(U(r"AI \u667a\u80fd\u4f53\u77e9\u9635"), '<div class="grid">' + cards + "</div>", user=user, wide=True))

    def business_overview(self, user):
        user = self.require_login(user)
        if not user:
            return
        s = load_summary()
        metrics = "".join([
            self.metric(U(r"\u672c\u6708\u9500\u552e"), U(r"\uffe5") + money(s.get("month_sales")), pct(s.get("completion_rate"))),
            self.metric(U(r"\u672c\u6708\u6bdb\u5229"), U(r"\uffe5") + money(s.get("month_gross_profit")), U(r"\u6bdb\u5229\u7387 ") + pct(s.get("yesterday_gross_margin"))),
            self.metric(U(r"\u5e93\u5b58\u91d1\u989d"), U(r"\uffe5") + money(s.get("inventory_amount")), U(r"\u98ce\u9669 ") + money(s.get("risk_count"))),
        ])
        body = f"<div class='panel'><h2>{U(r'\u7ecf\u8425\u603b\u89c8')}</h2><div class='metrics'>{metrics}</div>{self.bullets(s.get('ai_suggestions', []))}</div>"
        self.out(layout(U(r"\u7ecf\u8425\u603b\u89c8"), body, user=user, wide=True))

    def sap_sync(self, user):
        user = self.require_login(user)
        if not user:
            return
        s = load_summary()
        body = f"""
<div class="panel">
  <h2>SAP B1 {U(r'\u540c\u6b65\u72b6\u6001')}</h2>
  <div class="metrics">
    {self.metric(U(r'\u6570\u636e\u65e5\u671f'), s.get('data_date'), U(r'\u6458\u8981\u6587\u4ef6'))}
    {self.metric(U(r'\u6628\u65e5\u9500\u552e'), U(r'\uffe5') + money(s.get('yesterday_sales')), U(r'SAP B1'))}
    {self.metric(U(r'\u5e93\u5b58\u91d1\u989d'), U(r'\uffe5') + money(s.get('inventory_amount')), U(r'\u5df2\u540c\u6b65'))}
  </div>
  <p class="small">{U(r'\u624b\u52a8\u540c\u6b65\u5efa\u8bae\u5728\u670d\u52a1\u5668\u6267\u884c systemctl start firefox-sap-sync.service\uff0c\u9875\u9762\u5148\u63d0\u4f9b\u72b6\u6001\u67e5\u770b\u3002')}</p>
</div>"""
        self.out(layout(U(r"SAP B1 \u540c\u6b65"), body, user=user, wide=True))

    def dashboard(self, user):
        role = user["role"]
        can_boss = role in ("boss", "admin", "finance", "purchasing")
        can_manager = role in ("boss", "admin", "store_manager", "purchasing", "finance")
        can_admin = role == "admin"
        cards = [
            self.card(U(r"AI \u603b\u7ecf\u7406"), U(r"\u8001\u677f\u9a7e\u9a76\u8231\uff0c\u67e5\u770b\u7ecf\u8425\u5206\u6790\u548c\u51b3\u7b56\u5efa\u8bae\u3002"), "https://dify.huyan.vafox.com/chat/firefox-ai-ceo", "btn", can_boss),
            self.card(U(r"AI \u5e97\u957f"), U(r"\u95e8\u5e97\u65e5\u5e38\u7ba1\u7406\u3001\u8bdd\u672f\u3001\u4f1a\u5458\u8ddf\u8fdb\u548c\u4efb\u52a1\u63d0\u9192\u3002"), "https://dify.huyan.vafox.com/chat/firefox-ai-store-manager", "btn green", can_manager),
            self.card(U(r"\u706b\u72d0\u72f8\u77e5\u8bc6\u5e93"), U(r"\u67e5\u770b\u516c\u53f8\u67b6\u6784\u3001\u95e8\u5e97 SOP\u3001\u9879\u76ee\u548c\u57f9\u8bad\u8d44\u6599\u3002"), "/wiki/", "btn red", True),
            self.card(U(r"\u95e8\u5e97\u65e5\u62a5"), U(r"\u67e5\u770b\u6bcf\u65e5\u9500\u552e\u3001\u6bdb\u5229\u3001\u4efb\u52a1\u548c\u95e8\u5e97\u590d\u76d8\u3002"), "/wiki/firefox-hq/daily", "btn orange", True),
            self.card(U(r"\u5e93\u5b58\u5206\u6790"), U(r"\u67e5\u770b SAP B1 \u540c\u6b65\u72b6\u6001\u3001\u5e93\u5b58\u548c\u9500\u552e\u6570\u636e\u5165\u53e3\u3002"), "/wiki/firefox-hq/sap-b1", "btn dark", can_manager),
            self.card(U(r"\u7cfb\u7edf\u7ba1\u7406"), U(r"\u5ba1\u6838\u5458\u5de5\u3001\u7981\u7528\u8d26\u53f7\u3001\u4fee\u6539\u89d2\u8272\u548c\u91cd\u7f6e\u5bc6\u7801\u3002"), "/admin", "btn dark", can_admin),
        ]
        info = '<div class="panel"><strong>{}</strong><p class="small">{}：{} · {}：{} · {}：{}</p></div>'.format(
            U(r"\u5f53\u524d\u8d26\u53f7"),
            T["name"],
            esc(user["name"]),
            T["store"],
            esc(user["store"]),
            T["role"],
            esc(ROLES.get(role, role)),
        )
        self.out(layout(U(r"\u706b\u72d0\u72f8\u5de5\u4f5c\u53f0\u9996\u9875"), '<div class="grid">' + "".join(cards) + "</div>" + info, user=user))

    def dashboard(self, user):
        role = user["role"]
        can_boss = role in ("boss", "admin", "finance", "purchasing")
        can_manager = role in ("boss", "admin", "store_manager", "purchasing", "finance")
        can_admin = role == "admin"
        summary = load_summary()
        quick = f"""
<div class="panel">
  <h2>{U(r'\u5feb\u901f\u7ecf\u8425\u63d0\u793a')}</h2>
  <p class="small">{U(r'\u6628\u65e5\u9500\u552e')}：{U(r'\uffe5') + money(summary.get("yesterday_sales"))} ｜ {U(r'\u672c\u6708\u5b8c\u6210\u7387')}：{pct(summary.get("completion_rate"))} ｜ {U(r'\u6570\u636e\u65e5\u671f')}：{esc(summary.get("data_date"))}</p>
</div>"""
        cards = [
            self.card(U(r"AI \u603b\u7ecf\u7406"), U(r"\u6253\u5f00 AI \u603b\u7ecf\u7406\u6668\u62a5\uff0c\u67e5\u770b\u9500\u552e\u3001\u6bdb\u5229\u3001\u5e93\u5b58\u548c\u7ecf\u8425\u5efa\u8bae\u3002"), "/ai-ceo", "btn", can_boss),
            self.card(U(r"\u7ecf\u8425\u603b\u89c8"), U(r"\u8fdb\u5165\u9500\u552e\u3001\u6bdb\u5229\u3001\u5e93\u5b58\u548c\u98ce\u9669\u6570\u636e\u770b\u677f\u3002"), "/overview", "btn dark", can_boss),
            self.card(U(r"\u95e8\u5e97\u4e2d\u5fc3"), U(r"\u95e8\u5e97\u6863\u6848\u3001\u7ecf\u8425\u6570\u636e\u3001\u65f6\u95f4\u8f74\u548c AI \u5206\u6790\u3002"), "/stores", "btn green", can_manager),
            self.card(U(r"\u5458\u5de5\u4e2d\u5fc3"), U(r"\u5458\u5de5\u6863\u6848\u3001\u9500\u552e\u8868\u73b0\u3001\u57f9\u8bad\u8bb0\u5f55\u548c AI \u5efa\u8bae\u3002"), "/employees", "btn green", can_manager),
            self.card(U(r"\u54c1\u724c\u4e2d\u5fc3"), U(r"\u54c1\u724c\u6863\u6848\u3001\u5408\u4f5c\u6587\u4ef6\u3001\u54c1\u724c\u9500\u552e\u548c AI \u5206\u6790\u3002"), "/brands", "btn", role in ("boss", "admin", "purchasing", "store_manager")),
            self.card(U(r"\u4ea7\u54c1\u4e2d\u5fc3"), U(r"\u4ea7\u54c1\u6863\u6848\u3001\u5e93\u5b58\u3001\u9500\u552e\u6570\u636e\u548c AI \u8bdd\u672f\u3002"), "/products", "btn", True),
            self.card(U(r"\u4f9b\u5e94\u5546\u4e2d\u5fc3"), U(r"\u4f9b\u5e94\u5546\u6863\u6848\u3001\u5408\u540c\u3001\u91c7\u8d2d\u548c\u4ed8\u6b3e\u8bb0\u5f55\u3002"), "/suppliers", "btn orange", role in ("boss", "admin", "purchasing", "finance")),
            self.card(U(r"\u987e\u5ba2/\u4f1a\u5458\u4e2d\u5fc3"), U(r"\u4f1a\u5458\u6863\u6848\u3001\u8d2d\u4e70\u5386\u53f2\u3001\u504f\u597d\u6807\u7b7e\u548c\u7ef4\u62a4\u5efa\u8bae\u3002"), "/members", "btn green", True),
            self.card(U(r"\u8d22\u52a1\u4e2d\u5fc3"), U(r"\u5bf9\u516c\u8d26\u52a1\u3001\u73b0\u91d1\u8d26\u3001\u4ed8\u6b3e\u8bc4\u4f30\u548c\u8d44\u91d1\u8ba1\u5212\u3002"), "/finance", "btn dark", role in ("boss", "admin", "finance")),
            self.card(U(r"\u5e93\u5b58\u91c7\u8d2d\u4e2d\u5fc3"), U(r"\u67e5\u770b SAP B1 \u540c\u6b65\u3001\u5e93\u5b58\u91d1\u989d\u548c\u91c7\u8d2d\u9884\u8b66\u3002"), "/inventory", "btn dark", can_manager),
            self.card(U(r"\u5185\u5bb9\u53d1\u5e03\u4e2d\u5fc3"), U(r"\u65b0\u5a92\u4f53\u63a8\u5e7f\u3001\u95e8\u5e97\u5185\u5bb9\u3001\u4ea7\u54c1\u7d20\u6750\u548c\u53d1\u5e03\u8ba1\u5212\u3002"), "/content", "btn orange", True),
            self.card(U(r"\u77e5\u8bc6\u4e2d\u5fc3"), U(r"\u516c\u53f8\u5236\u5ea6\u3001SOP\u3001\u57f9\u8bad\u3001\u4ea7\u54c1\u8d44\u6599\u548c AI \u5f00\u53d1\u6587\u6863\u3002"), "/knowledge", "btn red", True),
            self.card(U(r"AI \u667a\u80fd\u4f53\u67e5\u8be2"), U(r"\u4f18\u5148\u67e5\u5185\u90e8\u77e5\u8bc6\u5e93\uff0c\u4e0d\u8db3\u65f6\u518d\u63a5\u5916\u7f51\u641c\u7d22\u3002"), "/ai-assistant", "btn", True),
            self.card(U(r"\u4efb\u52a1\u4e2d\u5fc3"), U(r"\u4eca\u65e5\u5f85\u529e\u3001\u95e8\u5e97\u4efb\u52a1\u3001\u81ea\u52a8\u5316\u4efb\u52a1\u548c\u8ddf\u8fdb\u63d0\u9192\u3002"), "/tasks", "btn", True),
            self.card(U(r"\u7cfb\u7edf\u7ba1\u7406"), U(r"\u5ba1\u6838\u5458\u5de5\u3001\u7981\u7528\u8d26\u53f7\u3001\u4fee\u6539\u89d2\u8272\u548c\u91cd\u7f6e\u5bc6\u7801\u3002"), "/admin", "btn dark", can_admin),
        ]
        info = '<div class="panel"><strong>{}</strong><p class="small">{}：{} ｜ {}：{} ｜ {}：{}</p></div>'.format(
            U(r"\u5f53\u524d\u8d26\u53f7"),
            T["name"],
            esc(user["name"]),
            T["store"],
            esc(user["store"]),
            T["role"],
            esc(ROLES.get(role, role)),
        )
        info = '<div class="panel"><strong>{}</strong><p class="small">{}：{} ｜ {}：{} ｜ {}：{}</p></div>'.format(
            U(r"\u5f53\u524d\u8d26\u53f7"), T["name"], esc(user["name"]), T["store"], esc(user["store"]), T["role"], esc(ROLES.get(role, role))
        )
        self.out(layout(T["brand"], quick + '<div class="grid">' + "".join(cards) + "</div>" + info, user=user, wide=True))

    def ai_ceo(self, user):
        user = self.require_login(user)
        if not user:
            return
        if not self.can_open(user, ("boss", "admin", "finance", "purchasing")):
            return self.dashboard(user)
        summary = load_summary()
        metrics = "".join(
            [
                self.metric(U(r"\u672c\u6708\u9500\u552e"), U(r"\uffe5") + money(summary.get("month_sales")), U(r"\u5b8c\u6210\u7387 ") + pct(summary.get("completion_rate"))),
                self.metric(U(r"\u672c\u6708\u6bdb\u5229"), U(r"\uffe5") + money(summary.get("month_gross_profit")), U(r"\u6628\u65e5\u6bdb\u5229\u7387 ") + pct(summary.get("yesterday_gross_margin"))),
                self.metric(U(r"\u5e93\u5b58\u91d1\u989d"), U(r"\uffe5") + money(summary.get("inventory_amount")), U(r"\u98ce\u9669\u6570\u91cf ") + money(summary.get("risk_count"))),
                self.metric(U(r"\u4f1a\u5458"), U(r"\u5f85\u63a5\u5165"), U(r"\u540e\u7eed\u4ece SAP/Dify \u8865\u5145")),
                self.metric(U(r"\u6570\u636e\u65e5\u671f"), summary.get("data_date"), U(r"\u6bcf\u65e5 2:00 \u81ea\u52a8\u540c\u6b65")),
            ]
        )
        body = f"""
<div class="panel">
  <h2>{U(r'AI \u603b\u7ecf\u7406\u6668\u62a5')}</h2>
  <div class="metrics">{metrics}</div>
  <div class="split">
    <div><h2>{U(r'\u7ecf\u8425\u5224\u65ad')}</h2>{self.bullets(summary.get("ai_suggestions", [])[:5])}</div>
    <div><h2>{U(r'\u4eca\u65e5\u52a8\u4f5c')}</h2>{self.bullets(summary.get("todos", []))}</div>
  </div>
  <p><a class="btn" href="https://dify.huyan.vafox.com">{U(r'\u7ee7\u7eed\u548c AI \u804a\u5929')}</a></p>
</div>"""
        self.out(layout(U(r"AI \u603b\u7ecf\u7406"), body, user=user, wide=True))

    def ai_store_manager(self, user):
        user = self.require_login(user)
        if not user:
            return
        if not self.can_open(user, ("boss", "admin", "store_manager", "purchasing", "finance")):
            return self.dashboard(user)
        summary = load_summary()
        stores = summary.get("top_stores", []) or []
        if user["role"] == "store_manager" and user["store"]:
            store_key = str(user["store"]).strip()
            stores = [s for s in stores if store_key in str(s.get("store", ""))] or stores[:1]
        store_rows = ""
        for item in stores[:8]:
            store_rows += '<div class="store-row"><strong>{}</strong><span>{}：￥{}</span><span>{}：￥{}</span></div>'.format(
                esc(item.get("store", U(r"\u672a\u547d\u540d\u95e8\u5e97"))),
                U(r"\u9500\u552e"),
                money(item.get("sales")),
                U(r"\u6bdb\u5229"),
                money(item.get("gross_profit")),
            )
        if not store_rows:
            store_rows = '<p class="small">{}</p>'.format(U(r"\u6682\u65e0\u95e8\u5e97\u6570\u636e\uff0c\u5148\u663e\u793a\u6d4b\u8bd5\u6a21\u677f\u3002"))
        body = f"""
<div class="panel">
  <h2>{U(r'AI \u5e97\u957f')}</h2>
  <p class="small">{U(r'\u5e97\u957f\u53ea\u770b\u81ea\u5df1\u6709\u6743\u9650\u7684\u95e8\u5e97\uff1b\u8001\u677f\u548c\u7ba1\u7406\u5458\u53ef\u770b\u5168\u90e8\u95e8\u5e97\u3002')}</p>
  {store_rows}
</div>
<div class="split">
  <div class="panel"><h2>{U(r'\u4eca\u65e5\u4efb\u52a1')}</h2>{self.bullets([U(r'\u68c0\u67e5\u672c\u5e97\u65e5\u76ee\u6807\u5dee\u989d\u3002'), U(r'\u8ddf\u8fdb\u91cd\u70b9\u4f1a\u5458\u56de\u8bbf\u3002'), U(r'\u76d8\u70b9\u9ad8\u5e93\u5b58\u548c\u4f4e\u9500\u552e\u5546\u54c1\u3002')])}</div>
  <div class="panel"><h2>{U(r'\u5e93\u5b58\u63d0\u9192')}</h2>{self.bullets([U(r'\u6ede\u9500\u548c\u5c3a\u7801\u7ed3\u6784\u5206\u6790\u5c06\u4ece SAP B1 \u7ee7\u7eed\u8865\u5145\u3002'), U(r'\u5f53\u524d\u5148\u7528\u5e93\u5b58\u91d1\u989d\u548c\u98ce\u9669\u6570\u91cf\u505a\u95e8\u5e97\u9884\u8b66\u3002')])}</div>
</div>"""
        self.out(layout(U(r"AI \u5e97\u957f"), body, user=user, wide=True))

    def knowledge(self, user):
        user = self.require_login(user)
        if not user:
            return
        cats = [
            (U(r"\u516c\u53f8\u5236\u5ea6"), "/wiki/firefox-hq/company-architecture"),
            ("SOP", "/wiki/firefox-hq/operations"),
            (U(r"\u54c1\u724c"), "/wiki/firefox-hq/brands"),
            (U(r"\u4ea7\u54c1"), "/wiki/firefox-hq/supply-chain"),
            (U(r"\u57f9\u8bad"), "/wiki/firefox-hq/tasks"),
            (U(r"\u5927\u5e97\u9879\u76ee"), "/wiki/firefox-hq/big-store-plan"),
            (U(r"AI \u5f00\u53d1\u6587\u6863"), "/wiki/firefox-hq/company-ai-server-plan"),
            (U(r"\u6587\u4ef6\u4e0a\u4f20"), "/upload"),
            (U(r"\u5916\u7f51\u641c\u7d22"), "/web-search"),
            (U(r"AI \u52a9\u624b"), "/ai-assistant"),
        ]
        pills = "".join('<a class="pill" href="{}">{}</a>'.format(esc(href), esc(name)) for name, href in cats)
        body = f"""
<div class="panel">
  <h2>{U(r'\u706b\u72d0\u72f8 AI \u4f01\u4e1a\u77e5\u8bc6\u5e93')}</h2>
  <p class="small">{U(r'\u6309\u7ecf\u8425\u573a\u666f\u5206\u7c7b\uff0c\u65b9\u4fbf\u5458\u5de5\u624b\u673a\u6253\u5f00\u3001\u641c\u7d22\u548c\u5b66\u4e60\u3002')}</p>
  <div>{pills}</div>
  <form method="get" action="/wiki/search" style="margin-top:14px">
    <label>{U(r'\u5168\u6587\u641c\u7d22')}</label><input name="q" placeholder="{U(r'\u8f93\u5165\u5173\u952e\u8bcd\uff0c\u5982\uff1a\u85aa\u916c\u3001\u5357\u5c71\u3001SAP B1')}">
    <p><button>{U(r'\u641c\u7d22\u77e5\u8bc6\u5e93')}</button></p>
  </form>
</div>"""
        self.out(layout(U(r"\u706b\u72d0\u72f8 AI \u4f01\u4e1a\u77e5\u8bc6\u5e93"), body, user=user))

    def inventory(self, user):
        user = self.require_login(user)
        if not user:
            return
        if not self.can_open(user, ("boss", "admin", "store_manager", "purchasing", "finance")):
            return self.dashboard(user)
        summary = load_summary()
        body = f"""
<div class="panel">
  <h2>{U(r'\u5e93\u5b58\u5206\u6790')}</h2>
  <div class="metrics">
    {self.metric(U(r'\u5e93\u5b58\u91d1\u989d'), U(r'\uffe5') + money(summary.get('inventory_amount')), U(r'SAP B1'))}
    {self.metric(U(r'\u98ce\u9669\u6570\u91cf'), money(summary.get('risk_count')), U(r'\u9700\u8ddf\u8fdb'))}
    {self.metric(U(r'\u6570\u636e\u65e5\u671f'), summary.get('data_date'), U(r'\u6bcf\u65e5 2:00'))}
  </div>
  <p><a class="btn dark" href="/wiki/firefox-hq/sap-b1">{U(r'\u67e5\u770b SAP B1 \u540c\u6b65\u8bf4\u660e')}</a></p>
</div>"""
        self.out(layout(U(r"\u5e93\u5b58\u5206\u6790"), body, user=user, wide=True))

    def change_password(self, user, msg=""):
        if not user:
            return self.redir("/login")
        body = f"""
<div class="panel form">
  <form method="post" action="/change-password">
    <label>{U(r'\u539f\u5bc6\u7801')}</label><input name="old_password" type="password" required>
    <label>{T['new_password']}</label><input name="new_password" type="password" required>
    <p><button>{T['save']}</button></p>
  </form>
</div>"""
        self.out(layout(T["change_password"], body, user=user, msg=msg))

    def change_password_post(self):
        user = self.current_user()
        if not user:
            return self.redir("/login")
        form = self.form()
        if not cp(form.get("old_password", ""), user["password_hash"]):
            return self.change_password(user, U(r"\u539f\u5bc6\u7801\u4e0d\u6b63\u786e\u3002"))
        with db() as conn:
            conn.execute(
                "update users set password_hash=?, reset_required=0, updated_at=? where id=?",
                (hp(form.get("new_password", "")), int(time.time()), user["id"]),
            )
        return self.change_password(user, T["password_changed"])

    def role_select(self, value):
        return '<select name="role">' + "".join(
            '<option value="{}"{}>{}</option>'.format(k, " selected" if value == k else "", esc(v)) for k, v in ROLES.items()
        ) + "</select>"

    def status_select(self, value):
        return '<select name="status">' + "".join(
            '<option value="{}"{}>{}</option>'.format(k, " selected" if value == k else "", esc(v)) for k, v in STATUS.items()
        ) + "</select>"

    def admin(self, msg=""):
        user = self.require_admin()
        if not user:
            return
        with db() as conn:
            rows = conn.execute("select * from users order by status='pending' desc, created_at desc").fetchall()
        body = f'<div class="panel"><h2>{T["users"]}</h2><p class="small">{U(r"\u65b0\u5458\u5de5\u6ce8\u518c\u540e\u9ed8\u8ba4\u4e3a\u5f85\u5ba1\u6838\u3002\u7ba1\u7406\u5458\u53ef\u4ee5\u5ba1\u6838\u901a\u8fc7\u3001\u7981\u7528\u8d26\u53f7\u3001\u4fee\u6539\u89d2\u8272\u548c\u91cd\u7f6e\u5bc6\u7801\u3002")}</p><table><thead><tr><th>{T["name"]}</th><th>{T["email"]}</th><th>{T["phone"]}</th><th>{T["store"]}</th><th>{T["role"]}</th><th>{T["status"]}</th><th>{T["action"]}</th></tr></thead><tbody>'
        for row in rows:
            update = (
                '<form method="post" action="/admin/update">'
                f'<input type="hidden" name="id" value="{row["id"]}">'
                f'{self.role_select(row["role"])}{self.status_select(row["status"])}'
                f'<button>{T["save"]}</button></form>'
            )
            reset = (
                '<form method="post" action="/admin/reset-password">'
                f'<input type="hidden" name="id" value="{row["id"]}">'
                f'<input name="new_password" type="text" placeholder="{T["new_password"]}" required>'
                f'<button class="gray">{T["reset"]}</button></form>'
            )
            body += "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td><div class=\"inline\">{}{}</div></td></tr>".format(
                esc(row["name"]),
                esc(row["email"]),
                esc(row["phone"]),
                esc(row["store"]),
                esc(ROLES.get(row["role"], row["role"])),
                esc(STATUS.get(row["status"], row["status"])),
                update,
                reset,
            )
        body += "</tbody></table></div>"
        self.out(layout(T["admin"], body, user=user, msg=msg, wide=True))

    def admin_update(self):
        user = self.require_admin()
        if not user:
            return
        form = self.form()
        role = form.get("role", "employee") if form.get("role") in ROLES else "employee"
        status = form.get("status", "pending") if form.get("status") in STATUS else "pending"
        with db() as conn:
            conn.execute("update users set role=?, status=?, updated_at=? where id=?", (role, status, int(time.time()), form.get("id")))
        return self.redir("/admin")

    def admin_reset_password(self):
        user = self.require_admin()
        if not user:
            return
        form = self.form()
        new_password = form.get("new_password", "")
        if len(new_password) < 6:
            return self.admin(U(r"\u65b0\u5bc6\u7801\u81f3\u5c11 6 \u4f4d\u3002"))
        with db() as conn:
            conn.execute(
                "update users set password_hash=?, failed_attempts=0, locked_until=0, reset_required=1, updated_at=? where id=?",
                (hp(new_password), int(time.time()), form.get("id")),
            )
        return self.redir("/admin")

    def dashboard(self, user):
        role = user["role"]
        can_boss = role in ("boss", "admin", "finance", "purchasing")
        can_manager = role in ("boss", "admin", "store_manager", "purchasing", "finance")
        can_admin = role == "admin"
        cards = [
            self.card(U(r"AI \u603b\u7ecf\u7406"), U(r"\u4eca\u5929\u516c\u53f8\u60c5\u51b5\u3001\u98ce\u9669\u63d0\u9192\u3001\u7ecf\u8425\u5efa\u8bae\u3002"), "/ai-ceo", "btn", can_boss),
            self.card(U(r"\u7ecf\u8425\u603b\u89c8"), U(r"\u9500\u552e\u3001\u5229\u6da6\u3001\u5e93\u5b58\u3001\u73b0\u91d1\u6d41\u3002"), "/business-overview", "btn dark", can_boss),
            self.card(U(r"\u95e8\u5e97\u4e2d\u5fc3"), U(r"\u5357\u5c71\u5e97\u3001\u632f\u5174\u5e97\u3001\u822a\u82d1\u5e97\u7b49\u95e8\u5e97\u6863\u6848\u3002"), "/stores", "btn green", can_manager),
            self.card(U(r"\u5458\u5de5\u4e2d\u5fc3"), U(r"\u5458\u5de5\u4fe1\u606f\u3001\u9500\u552e\u3001\u6536\u5165\u3001\u6210\u957f\u8bb0\u5f55\u3002"), "/employees", "btn green", can_manager),
            self.card(U(r"\u54c1\u724c\u4e2d\u5fc3"), U(r"KAILAS\u3001Mammut\u3001OSPREY\u3001Deuter \u7b49\u54c1\u724c\u6863\u6848\u3002"), "/brands", "btn", role in ("boss", "admin", "purchasing", "store_manager")),
            self.card(U(r"\u4ea7\u54c1\u4e2d\u5fc3"), U(r"\u4ea7\u54c1\u8bf4\u660e\u3001\u5e93\u5b58\u3001\u9500\u552e\u5386\u53f2\u3001AI \u8bdd\u672f\u3002"), "/products", "btn", True),
            self.card(U(r"\u4f9b\u5e94\u5546\u4e2d\u5fc3"), U(r"\u5408\u540c\u3001\u4ed8\u6b3e\u3001\u4f9b\u8d27\u8bb0\u5f55\u3002"), "/suppliers", "btn orange", role in ("boss", "admin", "purchasing", "finance")),
            self.card(U(r"\u987e\u5ba2\u4e2d\u5fc3"), U(r"\u6d88\u8d39\u8bb0\u5f55\u3001\u6807\u7b7e\u3001\u4f1a\u5458\u8d21\u732e\u3002"), "/members", "btn green", True),
            self.card(U(r"\u5e93\u5b58\u91c7\u8d2d"), U(r"\u5e93\u5b58\u98ce\u9669\u3001\u91c7\u8d2d\u9884\u8b66\u3001SAP B1 \u6458\u8981\u3002"), "/inventory", "btn dark", can_manager),
            self.card(U(r"\u8d22\u52a1\u4e2d\u5fc3"), U(r"\u8d44\u91d1\u3001\u4ed8\u6b3e\u3001\u8d39\u7528\u3001\u5229\u6da6\u5206\u6790\u3002"), "/finance", "btn dark", role in ("boss", "admin", "finance")),
            self.card(U(r"\u5185\u5bb9\u4e2d\u5fc3"), U(r"\u516c\u4f17\u53f7\u3001\u89c6\u9891\u53f7\u3001\u5c0f\u7ea2\u4e66\u3001\u6296\u97f3\u7b49\u5185\u5bb9\u6846\u67b6\u3002"), "/content-center", "btn orange", True),
            self.card(U(r"\u77e5\u8bc6\u4e2d\u5fc3"), U(r"\u6587\u6863\u3001\u5408\u540c\u3001\u4f1a\u8bae\u7eaa\u8981\u3001AI \u95ee\u7b54\u7684\u77e5\u8bc6\u5165\u53e3\u3002"), "/knowledge", "btn red", True),
            self.card(U(r"AI \u667a\u80fd\u4f53"), U(r"AI CEO\u3001AI CFO\u3001AI \u91c7\u8d2d\u3001AI \u5e93\u5b58\u7b49\u667a\u80fd\u4f53\u77e9\u9635\u3002"), "/agents", "btn", True),
            self.card(U(r"\u5de5\u4f5c\u6d41"), U(r"\u91c7\u8d2d\u3001\u4ed8\u6b3e\u3001\u5408\u540c\u3001\u8bf7\u5047\u3001\u5185\u5bb9\u5ba1\u6838\u6d41\u7a0b\u9884\u7559\u3002"), "/workflow", "btn", True),
            self.card(U(r"\u7cfb\u7edf\u7ba1\u7406"), U(r"\u8d26\u53f7\u3001\u6743\u9650\u3001\u5ba1\u6838\u548c\u91cd\u7f6e\u5bc6\u7801\u3002"), "/admin", "btn dark", can_admin),
        ]
        info = '<div class="panel"><strong>{}</strong><p class="small">{}：{} ｜ {}：{} ｜ {}：{}</p></div>'.format(
            U(r"\u5f53\u524d\u8d26\u53f7"), T["name"], esc(user["name"]), T["store"], esc(user["store"]), T["role"], esc(ROLES.get(role, role))
        )
        self.out(layout(T["brand"], '<div class="grid">' + "".join(cards) + "</div>" + info, user=user, wide=True))

    def module_page(self, user, path):
        user = self.require_login(user)
        if not user:
            return
        module = module_key(path)
        title, text, roles = MODULES[path]
        if not self.can_open(user, roles):
            return self.dashboard(user)
        q = parse_qs(urlparse(self.path).query).get("q", [""])[0].strip()
        sql = "select * from records where module=? and status!='deleted'"
        params = [module]
        if q:
            like = "%" + q + "%"
            sql += " and (title like ? or tags like ? or summary like ? or data_json like ?)"
            params += [like, like, like, like]
        sql += " order by updated_at desc limit 100"
        with db() as conn:
            rows = conn.execute(sql, params).fetchall()
        field_html = "".join('<span class="pill">{}</span>'.format(esc(f)) for f in ARCHIVE_FIELDS.get(module, []))
        row_html = ""
        for row in rows:
            row_html += "<tr><td>{}</td><td>{}</td><td>{}</td><td><a href=\"/records/view?id={}\">{}</a> <a href=\"/records/edit?id={}\">{}</a></td></tr>".format(
                esc(row["title"]), esc(row["tags"]), esc(dt(row["updated_at"])), row["id"], U(r"\u8be6\u60c5"), row["id"], U(r"\u7f16\u8f91")
            )
        if not row_html:
            row_html = '<tr><td colspan="4" class="small">{}</td></tr>'.format(U(r"\u6682\u65e0\u6863\u6848\uff0c\u53ef\u4ee5\u5148\u65b0\u5efa\u4e00\u6761\u3002"))
        abilities = [U(r"\u65b0\u5efa"), U(r"\u7f16\u8f91"), U(r"\u5220\u9664"), U(r"\u8be6\u60c5\u9875"), U(r"\u641c\u7d22"), U(r"\u6807\u7b7e"), U(r"\u5907\u6ce8"), U(r"\u65f6\u95f4\u8f74"), U(r"\u5173\u8054\u5bf9\u8c61"), U(r"\u9644\u4ef6"), U(r"\u56fe\u7247/PDF/Word/Excel/\u89c6\u9891\u4e0a\u4f20"), U(r"AI \u67e5\u8be2\u9884\u7559")]
        body = f"""
<div class="panel">
  <h2>{esc(title)}</h2><p>{esc(text)}</p>
  <div class="inline"><a class="btn" href="/records/new?module={esc(module)}">{U(r'\u65b0\u5efa')}</a><a class="btn green" href="/records/export?module={esc(module)}">{U(r'Excel \u5bfc\u51fa')}</a><a class="btn orange" href="/upload">{U(r'\u4e0a\u4f20\u9644\u4ef6')}</a><a class="btn gray" href="/ai-query">{U(r'AI \u67e5\u8be2')}</a></div>
</div>
<div class="panel"><h2>{U(r'\u5b57\u6bb5\u89c4\u5212')}</h2><div>{field_html}</div></div>
<div class="panel">
  <form method="get" action="/{esc(module)}"><label>{U(r'\u641c\u7d22')}</label><input name="q" value="{esc(q)}" placeholder="{U(r'\u641c\u7d22\u540d\u79f0\u3001\u6807\u7b7e\u3001\u5185\u5bb9')}"><p><button>{U(r'\u641c\u7d22')}</button></p></form>
  <table><thead><tr><th>{U(r'\u540d\u79f0')}</th><th>{U(r'\u6807\u7b7e')}</th><th>{U(r'\u66f4\u65b0')}</th><th>{T['action']}</th></tr></thead><tbody>{row_html}</tbody></table>
</div>
<div class="panel"><h2>{U(r'\u7edf\u4e00\u6863\u6848\u80fd\u529b')}</h2>{self.bullets(abilities)}</div>"""
        self.out(layout(title, body, user=user, wide=True))

    def document_center(self, user):
        user = self.require_login(user)
        if not user:
            return
        with db() as conn:
            files = conn.execute("select * from uploaded_files order by created_at desc limit 80").fetchall()
        rows = "".join("<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(esc(f["original_name"]), esc(f["category"]), esc(f["mime"]), esc(dt(f["created_at"]))) for f in files)
        if not rows:
            rows = '<tr><td colspan="4" class="small">{}</td></tr>'.format(U(r"\u6682\u65e0\u6587\u4ef6\uff0c\u8bf7\u5148\u4e0a\u4f20\u3002"))
        features = [U(r"\u4e0a\u4f20"), U(r"\u6587\u4ef6\u5217\u8868"), U(r"\u6587\u4ef6\u8be6\u60c5"), U(r"\u5173\u8054\u5bf9\u8c61"), U(r"\u6807\u7b7e"), U(r"OCR \u9884\u7559"), U(r"\u81ea\u52a8\u6458\u8981"), U(r"\u5411\u91cf\u5316\u9884\u7559"), U(r"AI Q&A")]
        body = f"<div class='panel'><h2>{U(r'\u7edf\u4e00\u6587\u4ef6\u4e2d\u5fc3')}</h2><p><a class='btn' href='/upload'>{U(r'\u4e0a\u4f20\u6587\u4ef6')}</a></p>{self.bullets(features)}</div><div class='panel'><table><thead><tr><th>{U(r'\u6587\u4ef6')}</th><th>{U(r'\u5206\u7c7b')}</th><th>{U(r'\u7c7b\u578b')}</th><th>{U(r'\u65f6\u95f4')}</th></tr></thead><tbody>{rows}</tbody></table></div>"
        self.out(layout(U(r"\u7edf\u4e00\u6587\u4ef6\u4e2d\u5fc3"), body, user=user, wide=True))

    def workflow_center(self, user):
        user = self.require_login(user)
        if not user:
            return
        flows = [U(r"\u91c7\u8d2d\u5ba1\u6279"), U(r"\u4ed8\u6b3e\u5ba1\u6279"), U(r"\u5408\u540c\u5ba1\u6279"), U(r"\u8bf7\u5047\u5ba1\u6279"), U(r"\u62db\u8058"), U(r"\u79bb\u804c"), U(r"\u95e8\u5e97\u4efb\u52a1\u5206\u914d"), U(r"\u5185\u5bb9\u5ba1\u6838")]
        cards = "".join(self.card(f, U(r"\u6d41\u7a0b\u5f15\u64ce\u9884\u7559\uff1a\u5ba1\u6279\u8282\u70b9\u3001\u6267\u884c\u4eba\u3001\u65f6\u95f4\u8f74\u3001\u901a\u77e5\u548c\u65e5\u5fd7\u3002"), "/tasks", "btn", True) for f in flows)
        self.out(layout(U(r"\u5de5\u4f5c\u6d41\u4e2d\u5fc3"), '<div class="grid">' + cards + "</div>", user=user, wide=True))

    def sap_api_placeholder(self, user, path):
        if not user:
            return self.json_out({"ok": False, "message": "login required"}, code=401)
        mapping = {
            "/api/sap/business-analysis": "business analysis",
            "/api/sap/profit-analysis": "profit analysis",
            "/api/sap/inventory-analysis": "inventory analysis",
            "/api/sap/sales-trend": "sales trend",
            "/api/sap/ai-analysis": "ai analysis",
        }
        return self.json_out({"ok": True, "endpoint": mapping.get(path, "sap placeholder"), "data": load_summary(), "note": "placeholder; existing SAP sync is unchanged"})

    def agents(self, user):
        user = self.require_login(user)
        if not user:
            return
        agent_defs = [
            (U(r"AI CEO / AI \u603b\u7ecf\u7406"), U(r"\u7ecf\u8425\u5224\u65ad\u3001\u98ce\u9669\u9884\u8b66\u3001\u884c\u52a8\u5efa\u8bae")),
            (U(r"AI CFO / AI \u8d22\u52a1\u603b\u76d1"), U(r"\u73b0\u91d1\u6d41\u3001\u8d39\u7528\u3001\u4ed8\u6b3e\u548c\u5229\u6da6\u5206\u6790")),
            (U(r"AI COO / AI \u8fd0\u8425\u603b\u76d1"), U(r"\u95e8\u5e97\u8fd0\u8425\u3001\u4efb\u52a1\u548c SOP \u63a8\u8fdb")),
            (U(r"AI \u91c7\u8d2d\u7ecf\u7406"), U(r"\u91c7\u8d2d\u3001\u8865\u8d27\u3001\u4f9b\u5e94\u5546\u548c\u8d26\u671f")),
            (U(r"AI \u5e93\u5b58\u7ecf\u7406"), U(r"\u5e93\u5b58\u91d1\u989d\u3001\u6ede\u9500\u3001\u8c03\u62e8\u548c\u98ce\u9669")),
            (U(r"AI \u54c1\u724c\u7ecf\u7406"), U(r"\u54c1\u724c\u9500\u552e\u3001\u6bdb\u5229\u3001\u5408\u4f5c\u548c\u672a\u6765\u89c4\u5212")),
            (U(r"AI \u95e8\u5e97\u7ecf\u7406"), U(r"\u95e8\u5e97\u9500\u552e\u3001\u4f1a\u5458\u3001\u5458\u5de5\u548c\u5e93\u5b58")),
            (U(r"AI \u8425\u9500\u7ecf\u7406"), U(r"\u5c0f\u7ea2\u4e66\u3001\u6296\u97f3\u3001\u89c6\u9891\u53f7\u3001\u516c\u4f17\u53f7")),
            (U(r"AI \u57f9\u8bad\u7ecf\u7406"), U(r"\u57f9\u8bad\u8bfe\u7a0b\u3001\u8003\u6838\u548c\u6210\u957f\u8bb0\u5f55")),
            (U(r"AI \u5ba2\u670d"), U(r"\u987e\u5ba2\u95ee\u9898\u3001\u552e\u540e\u548c\u4f1a\u5458\u7ef4\u62a4")),
            (U(r"AI \u79d8\u4e66"), U(r"\u4f1a\u8bae\u7eaa\u8981\u3001\u4efb\u52a1\u5206\u89e3\u548c\u63d0\u9192")),
        ]
        cards = "".join(self.card(name, desc + U(r"\u3002\u72b6\u6001\uff1a\u6846\u67b6\u5df2\u5c31\u7eea\uff0cAPI \u9884\u7559\u3002"), "/ai-query", "btn", True) for name, desc in agent_defs)
        self.out(layout(U(r"AI \u667a\u80fd\u4f53\u4e2d\u5fc3"), '<div class="grid">' + cards + "</div>", user=user, wide=True))


if __name__ == "__main__":
    ThreadingHTTPServer(("127.0.0.1", PORT), App).serve_forever()
